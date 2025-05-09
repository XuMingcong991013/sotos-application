import os
import re
import pandas as pd
import numpy as np
import requests
from datetime import datetime, timedelta
import tempfile
import shutil
import time
import logging
from openpyxl import load_workbook



wb = load_workbook('D:/sotos工作资料/202503 安道拓（重庆）/封华工时表-重庆安道拓-2025年03月.xlsx')
ws = wb.active


def excel_to_markdown_with_merged_cells(worksheet):
    """
    将Excel工作表（包括合并单元格）转换为Markdown格式
    """
    # 获取工作表尺寸
    max_row = worksheet.max_row
    max_col = worksheet.max_column
    
    # 创建一个矩阵来存储单元格值和跨度信息
    cell_matrix = [[None for _ in range(max_col + 1)] for _ in range(max_row + 1)]
    
    # 处理合并单元格
    merged_ranges = {}
    for merged_range in worksheet.merged_cells.ranges:
        min_row, min_col = merged_range.min_row, merged_range.min_col
        max_row_range, max_col_range = merged_range.max_row, merged_range.max_col
        
        # 获取合并区域的值（使用左上角单元格的值）
        value = worksheet.cell(min_row, min_col).value
        
        # 记录这个合并区域
        merged_ranges[(min_row, min_col)] = (max_row_range - min_row + 1, max_col_range - min_col + 1)
        
        # 将值分配给合并区域内的所有单元格，但标记非主单元格
        for r in range(min_row, max_row_range + 1):
            for c in range(min_col, max_col_range + 1):
                if r == min_row and c == min_col:
                    # 主单元格（左上角）
                    cell_matrix[r][c] = (value, True, max_row_range - min_row + 1, max_col_range - min_col + 1)
                else:
                    # 非主单元格，标记为被合并
                    cell_matrix[r][c] = (None, False, 0, 0)
    
    # 填充非合并单元格的值
    for row in range(1, max_row + 1):
        for col in range(1, max_col + 1):
            if cell_matrix[row][col] is None:  # 如果不是合并单元格的一部分
                value = worksheet.cell(row, col).value
                cell_matrix[row][col] = (value, True, 1, 1)  # 正常单元格，跨度为1x1
    
    # 创建Markdown表格
    markdown_lines = []
    
    # 第一行：表格头部
    header_line = "| "
    for col in range(1, max_col + 1):
        header_line += "   | "
    markdown_lines.append(header_line)
    
    # 第二行：分隔符
    separator_line = "| "
    for col in range(1, max_col + 1):
        separator_line += "--- | "
    markdown_lines.append(separator_line)
    
    # 数据行
    for row in range(1, max_row + 1):
        row_line = "| "
        for col in range(1, max_col + 1):
            cell_info = cell_matrix[row][col]
            
            if cell_info[1]:  # 如果是主单元格
                value = cell_info[0]
                rowspan = cell_info[2]
                colspan = cell_info[3]
                
                # 转换为字符串，处理None值
                value_str = str(value) if value is not None else ""
                
                # 如果是合并单元格，添加标记
                if rowspan > 1 or colspan > 1:
                    value_str += f" [合并:{rowspan}x{colspan}]"
                
                row_line += value_str + " | "
            else:
                # 被合并的单元格，使用特殊标记
                row_line += "↑↑ | "
        
        markdown_lines.append(row_line)
    
    # 添加合并单元格的注释说明
    if merged_ranges:
        markdown_lines.append("\n**合并单元格说明:**")
        for (r, c), (rowspan, colspan) in merged_ranges.items():
            value = worksheet.cell(r, c).value
            value_str = str(value) if value is not None else "空值"
            markdown_lines.append(f"- 单元格 ({r},{c}) 合并范围 {rowspan}x{colspan}, 值: {value_str}")
    
    return "\n".join(markdown_lines)



def excel_to_html_markdown(worksheet):
    """
    将Excel工作表转换为带HTML表格的Markdown
    这种方法可以完美表示合并单元格
    """
    # 获取合并单元格信息
    merged_cells = {}
    for merged_range in worksheet.merged_cells.ranges:
        for row in range(merged_range.min_row, merged_range.max_row + 1):
            for col in range(merged_range.min_col, merged_range.max_col + 1):
                merged_cells[(row, col)] = (
                    merged_range.min_row,
                    merged_range.min_col,
                    merged_range.max_row - merged_range.min_row + 1,
                    merged_range.max_col - merged_range.min_col + 1
                )
    
    # 开始构建HTML表格
    html_lines = ['<table border="1">']
    
    # 添加表格行
    for row in range(1, worksheet.max_row + 1):
        html_lines.append('  <tr>')
        
        col = 1
        while col <= worksheet.max_column:
            # 检查当前单元格是否在合并单元格中
            if (row, col) in merged_cells:
                min_row, min_col, rowspan, colspan = merged_cells[(row, col)]
                
                # 只为合并区域的左上角单元格创建TD标签
                if row == min_row and col == min_col:
                    cell = worksheet.cell(row, col)
                    value = cell.value if cell.value is not None else ""
                    
                    html_lines.append(f'    <td rowspan="{rowspan}" colspan="{colspan}">{value}</td>')
                
                # 跳过合并区域内的其他单元格
                col += 1
            else:
                # 普通单元格
                cell = worksheet.cell(row, col)
                value = cell.value if cell.value is not None else ""
                html_lines.append(f'    <td>{value}</td>')
                col += 1
        
        html_lines.append('  </tr>')
    
    html_lines.append('</table>')
    
    # 将HTML表格嵌入到Markdown中
    markdown = "## Excel表格（包含合并单元格）\n\n"
    markdown += "\n".join(html_lines)
    
    return markdown



def excel_to_visual_markdown(worksheet):
    """
    将Excel工作表转换为视觉增强的Markdown表格
    使用边框字符表示合并单元格
    """
    max_row = worksheet.max_row
    max_col = worksheet.max_column
    
    # 创建一个空的字符矩阵来表示表格
    # 每个单元格占用多个字符位置以便表示边框
    cell_width = 12  # 每个单元格的字符宽度
    table = [[' ' for _ in range(max_col * (cell_width + 1) + 1)] for _ in range(max_row * 2 + 1)]
    
    # 绘制表格基础网格
    for row in range(max_row + 1):
        for col in range(max_col * (cell_width + 1) + 1):
            table[row * 2][col] = '-'  # 水平线
    
    for col in range(max_col + 1):
        for row in range(max_row * 2 + 1):
            table[row][col * (cell_width + 1)] = '|'  # 垂直线
    
    # 绘制交叉点
    for row in range(max_row + 1):
        for col in range(max_col + 1):
            table[row * 2][col * (cell_width + 1)] = '+'
    
    # 处理合并单元格和填充内容
    merged_ranges = {}
    for merged_range in worksheet.merged_cells.ranges:
        min_row, min_col = merged_range.min_row, merged_range.min_col
        max_row_range, max_col_range = merged_range.max_row, merged_range.max_col
        
        # 记录合并范围
        merged_ranges[(min_row, min_col)] = (max_row_range, max_col_range)
        
        # 移除合并范围内的内部网格线
        for r in range(min_row, max_row_range + 1):
            for c in range(min_col, max_col_range + 1):
                # 移除水平内部线
                if r < max_row_range and c <= max_col_range:
                    for x in range(cell_width + 1):
                        if c < max_col_range or x < cell_width:
                            table[r * 2][c * (cell_width + 1) + x] = ' '
                
                # 移除垂直内部线
                if c < max_col_range and r <= max_row_range:
                    for y in range(2):
                        table[r * 2 - 1 + y][c * (cell_width + 1)] = ' '
    
    # 填充单元格内容
    for row in range(1, max_row + 1):
        for col in range(1, max_col + 1):
            # 检查是否是合并单元格的主单元格
            is_main_cell = True
            for (min_row, min_col), (max_row_range, max_col_range) in merged_ranges.items():
                if row >= min_row and row <= max_row_range and col >= min_col and col <= max_col_range:
                    if row != min_row or col != min_col:
                        is_main_cell = False
                        break
            
            if is_main_cell:
                # 获取单元格值
                cell_value = worksheet.cell(row, col).value
                value_str = str(cell_value) if cell_value is not None else ""
                
                # 截断过长的值
                if len(value_str) > cell_width - 2:
                    value_str = value_str[:cell_width - 5] + "..."
                
                # 将值放入表格
                for i, char in enumerate(value_str):
                    table[row * 2 - 1][col * (cell_width + 1) - cell_width + i + 1] = char
    
    # 转换表格矩阵为字符串
    result = []
    for row in table:
        result.append(''.join(row))
    
    return "```\n" + "\n".join(result) + "\n```"


def excel_to_best_markdown(worksheet, prefer_html=False):
    """
    将Excel工作表转换为最佳Markdown表示
    
    Args:
        worksheet: openpyxl工作表对象
        prefer_html: 是否优先使用HTML表示（更准确但有些平台可能不支持）
    
    Returns:
        str: Markdown格式的表格表示
    """
    # 检查是否有合并单元格
    has_merged_cells = len(worksheet.merged_cells.ranges) > 0
    
    if has_merged_cells:
        if prefer_html:
            # 使用HTML表示合并单元格（最准确）
            return excel_to_html_markdown(worksheet)
        else:
            # 使用带注释的增强Markdown表示
            return excel_to_markdown_with_merged_cells(worksheet)
    else:
        # 没有合并单元格，使用标准Markdown表格
        return excel_to_standard_markdown(worksheet)

def excel_to_standard_markdown(worksheet):
    """将Excel工作表转换为标准Markdown表格（无合并单元格）"""
    markdown_lines = []
    max_row = worksheet.max_row
    max_col = worksheet.max_column
    
    # 第一行数据作为表头
    header_line = "| "
    separator_line = "| "
    
    for col in range(1, max_col + 1):
        cell = worksheet.cell(1, col)
        value = str(cell.value) if cell.value is not None else ""
        header_line += value + " | "
        separator_line += "--- | "
    
    markdown_lines.append(header_line)
    markdown_lines.append(separator_line)
    
    # 数据行
    for row in range(2, max_row + 1):
        row_line = "| "
        for col in range(1, max_col + 1):
            cell = worksheet.cell(row, col)
            value = str(cell.value) if cell.value is not None else ""
            row_line += value + " | "
        markdown_lines.append(row_line)
    
    return "\n".join(markdown_lines)


markdown_text = excel_to_markdown_with_merged_cells(ws)


# html_markdown = excel_to_html_markdown(ws)
# print(html_markdown)

def get_completions(user_prompt, system_prompt, stream=False, temperature=0.95, top_p=0.1):
    """请求模型，获取响应
    入参：
        - user_prompt: 用户输入
        - system_prompt: 身份设定
        - stream: 是否使用流式返回
        - temperature: (0,1], 控制生成文本的随机性
        - top_p: (0,1], 模型解码器只考虑从前top_p的概率的候选集中取tokens
    返回：
        - response: 模型返回的完整响应
    """
    url = "http://192.168.1.222:10000/openapi/community/v1/chat/completions"
    headers = {
        "Content-Type": "application/json",
        "Authorization": "Bearer ak_e89e4b0f5447234dfab930d2fc3790bfda1102bf47aaf17e9a723daa037a419c"
    }
    payload = {
        "user":"sotos001",
        "model":"QWen2.5-7B-Instruct",
        "messages":[
            {"role":"system", "content":system_prompt},
            {"role": "user", "content": user_prompt}  
        ],
        "stream":stream,
        "temperature": temperature,
        "top_p": top_p,
    }

    try:
        logging.info(f"发送请求到模型API...")
        logging.info(f"system_prompt：{system_prompt}")
        logging.info(f"user_prompt：{user_prompt}")
        response = requests.post(url, json=payload, headers=headers)
        response.raise_for_status()  # 如果请求失败会抛出HTTPError异常
        logging.info(f"请求成功，状态码: {response.status_code}")
        response = response.json()
        logging.info("模型完整响应内容: %s", response)

        return response
    
    except requests.exceptions.HTTPError as errh:
        logging.error(f"HTTP错误: {errh}\n请求URL: {url}\n请求负载: {payload}")
    except requests.exceptions.ConnectionError as errc:
        logging.error(f"连接错误: {errc}\n请检查网络连接和服务地址: {url}")
    except requests.exceptions.Timeout as errt:
        logging.error(f"请求超时: {errt}\n考虑增加超时时间或重试")
    except requests.exceptions.RequestException as err:
        logging.error(f"请求异常: {err}")
    except Exception as e:
        logging.error(f"未处理的异常: {e}", exc_info=True)



user_prompt = markdown_text
print(f"user_prompt:\n{user_prompt}\n\n\n")
# system_prompt = """
#     你是一个专业的工时数据分析师，你需要根据用户提供的工时表格，识别出三种主要区域：
#     1. 元数据区域：包含员工姓名、公司信息等基础信息，通常在表格顶部
#     2. 日常工时记录区域：包含日期、项目、工时时长等每日记录
#     3. 汇总区域：包含按项目汇总的工时数据，通常在表格底部

#     请注意区域之间没有明显分隔，需要通过内容语义和数据特征来识别。
#     请注意，忽略表格中的备注列。

#     请回答以下问题：
#     1. 员工姓名`employee_name`是什么？
#     2. 公司名称`company_name`是什么？
#     3. 如果表格中每一天的员工加班和工时信息，请严格按照项目名称做分组，分组计算出当月员工加班的小时数`overtime_hours`和总工时`total_hours`。形成一个新的表格，包含项目名称`project_name`、加班小时数`overtime_hours`和总工时`total_hours`。
#     注：请忽略表格中的所有备注信息，项目名称在表格中有显式的列，不要另外提取。
#     4. 如果表格中有关于工时按照项目分类的汇总信息，那么提取出汇总信息，与步骤3的计算结果做比较，判断是否一致。判断结果`compare_result`为`True`或`False`。

#     请以JSON格式返回,示例返回：
#     {“employee_name”:"张三", “company_name”:"AAA公司", “compare_result”:"True", 
#     "project_data": [{"project_name":"A11", "overtime_hours":2, "total_hours": 200}, {"project_name":"B11", "overtime_hours":1, "total_hours": 100}] }

    
# """

system_prompt_1 = """
你是一个专业的工时数据分析师。用户会提供给你一份Excel工时表，已转换为Markdown格式。
表格可能包含三个部分：（不一定每个部分都有）
1. 顶部元数据区域：包含公司名称、员工姓名等信息
2. 中间工时记录区域：每日详细工时记录
3. 底部汇总区域：按项目统计的工时汇总

请分析用户提供的数据并回答问题：
1. 表格中员工的姓名是什么？
2. 公司名称是什么？
3. 本月工时总计是多少小时？
4. 本月加班时长总计是多少小时？
5. 请按项目分别统计工时和加班时长。
6. 如果表格中有工时汇总数据，汇总数据显示加班时长多少小时、工时多少小时？
7. 请比对您的计算结果是否与表格汇总一致。

请提供详细的分析过程和结果。请一步步做分析,展示计算的每一步过程，加法不仅要展示sum公式，还要请展示具体的数值。
    
"""

response_1 = get_completions(user_prompt, system_prompt=system_prompt_1)
response_1 = response_1["choices"][0]["message"]["content"]
print(f"第一个答案：\n{response_1}\n##############################################################################")


system_prompt_2 = """
    你是一个专业的工时数据分析师，用户会提供给你一份另一个AI助手的回答，请从这份答案中提取出以下信息：
    1. 员工姓名`employee_name`是什么？
    2. 公司名称`company_name`是什么？
    3. 按照项目名称做分组，分组计算出当月员工加班的小时数`overtime_hours`和总工时`total_hours`。形成一个新的表格，包含项目名称`project_name`、加班小时数`overtime_hours`和总工时`total_hours`。
    4. 如果有汇总数据和明细数据的比较结果，请找出比较结果是一致还是不一致。一致是True，不一致是False。判断结果`compare_result`为`True`或`False`。

    请以JSON格式返回,示例返回：
    {“employee_name”:"张三", “company_name”:"AAA公司", “compare_result”:"True", 
    "project_data": [{"project_name":"A11", "overtime_hours":2, "total_hours": 200}, {"project_name":"B11", "overtime_hours":1, "total_hours": 100}] }

    注意：直接返回结果，不要包含任何解释或说明。
"""


user_prompt_2 = response_1
response_2 = get_completions(user_prompt=user_prompt_2, system_prompt=system_prompt_2)
response_2 = response_2["choices"][0]["message"]["content"]
# print(response_2)
print(f"第2个答案：\n{response_2}\n##############################################################################")

system_prompt_3 = """
你是一个专业的工时数据分析师，用户会提供给你一份Excel工时表，已转换为Markdown格式，还有一份关于数据的统计信息。请校验统计的数据是否正确。如果正确，输出True，否则输出False。

表格可能包含三个部分：（不一定每个部分都有）
1. 顶部元数据区域：包含公司名称、员工姓名等信息
2. 中间工时记录区域：每日详细工时记录
3. 底部汇总区域：按项目统计的工时汇总
    
统计数据：
- `employee_name`是员工姓名。
- `company_name`是公司名称。
- `compare_result`是原表格中工时明细数据与工时汇总数据的比较结果，是True或False。
- `project_data`是一个列表，包含项目名称、加班小时数和总工时，是工时明细数据根据项目名称分组计算的结果。

注意：直接返回结果，不要包含任何解释或说明。
"""

user_prompt_3 = f"""
统计数据：{response_2}

Excel工时表：
{markdown_text}
"""


response_3 = get_completions(user_prompt=user_prompt_3, system_prompt=system_prompt_3)
response_3 = response_3["choices"][0]["message"]["content"]
print(f"第3个答案：\n{response_3}\n##############################################################################")



#   注意：直接返回结果，不要包含任何解释或说明。






