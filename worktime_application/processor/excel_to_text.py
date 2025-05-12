import openpyxl
from openpyxl.utils import get_column_letter, range_boundaries
from datetime import datetime, timedelta
import re
import os


# 1.特殊处理Excel表格中的日期格式
def is_date_format(number_format):
    """
    检查一个数字格式是否是日期格式
    """
    date_formats = [
        'yyyy', 'yy', 'mm', 'dd', 'm/d', 'd/m',
        'mmm', 'mmmm', 'dddd',
        'yyyy-mm-dd', 'yy-mm-dd', 'mm-dd-yy', 'mm-dd-yyyy',
        'm/d/yy', 'd/m/yy', 'm/d/yyyy', 'd/m/yyyy',
        'yyyy/mm/dd', 'yy/mm/dd', 'mm/dd/yy', 'mm/dd/yyyy',
        'h:mm', 'h:mm:ss', 'h时', 'm分', 's秒',
        'yyyy年m月d日', 'yyyy年mm月dd日', 'm月d日', 'mm月dd日'
    ]
    
    # 将格式字符串转为小写进行比较
    number_format_lower = number_format.lower()
    
    for format_str in date_formats:
        if format_str.lower() in number_format_lower:
            return True
    
    return False

def is_likely_date(value, column, number_format):
    """
    检查一个数值是否可能是日期，考虑列位置和格式
    
    Parameters:
    - value: 单元格的值
    - column: 列索引 (0-based)
    - number_format: 单元格的数字格式
    """
    # 基本检查：值必须是数字且在合理范围内
    if not isinstance(value, (int, float)) or value < 1 or value > 73000:
        return False
    
    # 检查1：如果有明确的日期格式，则认为是日期
    if is_date_format(number_format):
        return True
    
    # 检查2：位置检查 - 日期更可能出现在前几列，特别是A列
    # 如果是A列(column=0)，使用相对宽松的标准
    if column == 0:
        # A列更可能是日期，但仍需要一些基本检查
        # 典型的Excel日期值范围（从1900-01-01到2100左右）
        return 1 <= value <= 73000
    
    # 检查3：对于非A列，使用更严格的标准
    # 非A列的数字很可能不是日期，除非有明确的日期格式或其他强有力的证据
    # 默认非A列的数字不是日期
    return False

def excel_date_to_datetime(excel_date):
    """
    将Excel日期值转换为Python datetime
    处理1900年闰年错误
    """
    if excel_date < 60:
        # 1900年1月和2月的日期
        return datetime(1899, 12, 31) + timedelta(days=excel_date)
    elif excel_date == 60:
        # 特殊情况：Excel错误地认为1900年2月29日存在
        return datetime(1900, 2, 28)
    else:
        # 其他日期需要减去1天来校正错误
        return datetime(1899, 12, 30) + timedelta(days=excel_date)

def format_cell_value(cell, column_index):
    """
    格式化单元格值，尤其是日期
    使用多种方法检测和处理日期，考虑列位置
    
    Parameters:
    - cell: 单元格对象
    - column_index: 列索引 (0-based)
    """
    if cell.value is None:
        return ""
    
    # 获取单元格的数字格式
    number_format = cell.number_format if hasattr(cell, 'number_format') else ""
    
    # 方法1：检查是否已经是datetime对象
    if isinstance(cell.value, datetime):
        # 根据单元格格式确定如何显示日期
        if '时' in number_format or ':' in number_format:
            return cell.value.strftime('%Y-%m-%d %H:%M:%S')
        else:
            return cell.value.strftime('%Y-%m-%d')
    
    # 方法2：通过数字格式、值和列位置判断是否是日期
    is_date_by_format = is_date_format(number_format)
    is_date_by_value_and_position = is_likely_date(cell.value, column_index, number_format)
    
    # 如果是数字且被判断为日期
    if isinstance(cell.value, (int, float)) and (is_date_by_format or is_date_by_value_and_position):
        try:
            date_obj = excel_date_to_datetime(cell.value)
            
            # 根据单元格格式确定如何显示日期
            if '时' in number_format or ':' in number_format:
                return date_obj.strftime('%Y-%m-%d %H:%M:%S')
            else:
                return date_obj.strftime('%Y-%m-%d')
        except (ValueError, OverflowError) as e:
            # 如果转换失败，记录错误并返回原值
            print(f"日期转换错误: {cell.value} ({e})")
            return str(cell.value)
    
    # 处理数字格式
    if isinstance(cell.value, (int, float)):
        # 检查是否有自定义数字格式
        if number_format and number_format != 'General':
            try:
                # 处理百分比格式
                if '%' in number_format:
                    return f"{cell.value:.2%}"
                # 处理带小数点的格式
                elif '0.00' in number_format:
                    return f"{cell.value:.2f}"
                elif '0.0' in number_format:
                    return f"{cell.value:.1f}"
            except:
                pass
    
    # 默认返回字符串值
    return str(cell.value)


# 2.表格转为markdown
def excel_to_markdown(file_path):
    """
    将Excel文件直接转换为Markdown格式，处理合并单元格和日期
    增强了日期检测和转换，考虑列位置
    """
    # 加载工作簿，确保数据被读取而不是公式
    wb = openpyxl.load_workbook(file_path, data_only=True)
    
    all_markdown_tables = []
    
    # 处理每个工作表
    for sheet_name in wb.sheetnames:
        sheet = wb[sheet_name]
        print(f"处理工作表: {sheet_name}")
        
        # 获取工作表的有效数据范围
        max_row = sheet.max_row
        max_col = sheet.max_column
        
        # 创建一个矩阵来存储扩展后的表格数据（处理合并单元格）
        expanded_data = [[None for _ in range(max_col)] for _ in range(max_row)]
        
        # 记录单元格类型，用于调试
        cell_types = [[None for _ in range(max_col)] for _ in range(max_row)]
        cell_formats = [[None for _ in range(max_col)] for _ in range(max_row)]
        
        # 首先填充所有非合并单元格
        for row in range(1, max_row + 1):
            for col in range(1, max_col + 1):
                cell = sheet.cell(row=row, column=col)
                
                # 记录单元格类型和格式，用于调试
                cell_types[row-1][col-1] = type(cell.value).__name__ if cell.value is not None else "None"
                cell_formats[row-1][col-1] = cell.number_format if hasattr(cell, 'number_format') else "Unknown"
                
                # 格式化单元格值，特别是日期，考虑列位置
                expanded_data[row-1][col-1] = format_cell_value(cell, col-1)
        
        # 记录一些调试信息，帮助识别日期问题
        debug_info = "## 调试信息（前5行5列）\n\n"
        debug_info += "### 单元格值类型\n\n"
        for row in range(min(5, max_row)):
            row_str = "| "
            for col in range(min(5, max_col)):
                row_str += f"{cell_types[row][col]} | "
            debug_info += row_str + "\n"
        
        debug_info += "\n### 单元格格式\n\n"
        for row in range(min(5, max_row)):
            row_str = "| "
            for col in range(min(5, max_col)):
                row_str += f"{cell_formats[row][col]} | "
            debug_info += row_str + "\n"
        
        debug_info += "\n\n"
        
        # 处理合并单元格
        for merged_range in sheet.merged_cells.ranges:
            min_col, min_row, max_col_range, max_row_range = range_boundaries(str(merged_range))
            top_left_cell = sheet.cell(row=min_row, column=min_col)
            top_left_value = expanded_data[min_row-1][min_col-1]  # 使用已格式化的值
            
            # 标记合并单元格
            for row in range(min_row, max_row_range + 1):
                for col in range(min_col, max_col_range + 1):
                    if row == min_row and col == min_col:
                        # 第一个单元格保持其值
                        expanded_data[row-1][col-1] = top_left_value
                    else:
                        # 使用符号标记其他合并单元格
                        direction = ""
                        if row > min_row:
                            direction += "^"  # 向上合并
                        if col > min_col:
                            direction += "<-"  # 向左合并
                        expanded_data[row-1][col-1] = direction
        
        # 确定实际有数据的范围，去除尾部空行空列
        actual_max_row = 0
        actual_max_col = 0
        
        for row in range(max_row):
            for col in range(max_col):
                if expanded_data[row][col] not in [None, ""]:
                    actual_max_row = max(actual_max_row, row + 1)
                    actual_max_col = max(actual_max_col, col + 1)
        
        # 构建markdown表格
        markdown_table = f"## 工作表: {sheet_name}\n\n"
        
        # 添加表头行
        header_row = "| "
        for col in range(1, actual_max_col + 1):
            column_letter = get_column_letter(col)
            header_row += f"{column_letter} | "
        markdown_table += header_row + "\n"
        
        # 添加分隔行
        separator = "| " + " --- |" * actual_max_col
        markdown_table += separator + "\n"
        
        # 添加数据行
        for row in range(actual_max_row):
            if all(expanded_data[row][col] in [None, ""] for col in range(actual_max_col)):
                continue  # 跳过完全空的行
            
            row_text = "| "
            for col in range(actual_max_col):
                value = expanded_data[row][col] if expanded_data[row][col] is not None else ""
                row_text += f"{value} | "
            
            markdown_table += row_text + "\n"
        
        # 添加调试信息
        # markdown_table += debug_info
        
        all_markdown_tables.append(markdown_table + "\n\n")
    
    return "".join(all_markdown_tables)




# test:
# file_path = "D:/test/于浩鑫-工时表-重庆安道拓-2025年03月.xlsx"
# markdown_text = excel_to_markdown(file_path)
# print(markdown_text)
