"""
根据结构化简历JSON生成标准Word简历。

最终Word属于交付结果，写入output_dir/final_resumes；解析Markdown、
结构化JSON、Excel和错误日志仍保留在output_dir/process_data。
"""

from __future__ import annotations

import json
import math
import shutil
import unicodedata
from copy import deepcopy
from pathlib import Path
from typing import Any
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import CONTENT_TYPE, RELATIONSHIP_TYPE
from docx.opc.packuri import PackURI
from docx.parts.numbering import NumberingPart
from docx.shared import Cm, Pt, RGBColor
from docx.text.run import Run

from utils.error_logging import write_error_log
from utils.processing_records import (
    TRACKING_WORKBOOK_NAME,
    update_generation_record,
)
from utils.runtime_paths import resource_path
from .supplement_report import (
    SUPPLEMENT_WORKBOOK_NAME,
    update_supplement_report,
)
from .markdown_renderer import (
    MarkdownBlock,
    find_markdown_residues,
    parse_resume_markdown,
)


FINAL_RESUMES_DIRNAME = "final_resumes"
PROCESS_DATA_DIRNAME = "process_data"
TEMPLATE_PATH = resource_path("templates", "standard_resume.docx")

FONT_NAME = "微软雅黑"
ACCENT_COLOR = RGBColor(84, 141, 212)
BODY_COLOR = RGBColor(0, 0, 0)
PLACEHOLDER_COLOR = RGBColor(255, 0, 0)
SMALL_FOUR = Pt(12)
FIVE = Pt(10.5)
BODY_LINE_SPACING = 1.0
SECTION_CONTENT_SPACE_BEFORE = Pt(12)
WORK_ENTRY_SPACE_BEFORE = Pt(10.5)
PROJECT_CONTENT_SPACE_BEFORE = Pt(18)
PLACEHOLDER = "待补充"

SECTION_LABELS = (
    "基本资料",
    "专业技能",
    "工作经历",
    "项目经验",
    "教育经历",
    "工作年限",
)


def ResumeGenerator(
    input_file: str,
    output_dir: str,
    with_photo: bool = False,
) -> str | None:
    """
    将InformationExtractor生成的JSON写入标准简历模板。

    Args:
        input_file: 结构化提取JSON文件路径。
        output_dir: 整个任务的输出根目录。
        with_photo: 是否在基本资料右侧保留证件照占位。

    Returns:
        成功时返回final_resumes目录下最终DOCX的绝对路径；
        真正异常写入过程数据错误日志并返回None。
    """

    return _resume_generator(
        input_file=input_file,
        output_dir=output_dir,
        with_photo=with_photo,
        remove_headers=False,
    )


def YouzuResumeGenerator(
    input_file: str,
    output_dir: str,
) -> str | None:
    """使用无照片布局生成不含页眉的优族标准简历。"""

    return _resume_generator(
        input_file=input_file,
        output_dir=output_dir,
        with_photo=False,
        remove_headers=True,
    )


def _resume_generator(
    input_file: str,
    output_dir: str,
    with_photo: bool,
    remove_headers: bool,
) -> str | None:
    """执行共用的标准简历生成、记录更新和错误处理流程。"""

    json_path = Path(input_file)
    output_root = Path(output_dir)
    process_directory = output_root / PROCESS_DATA_DIRNAME
    tracking_workbook = process_directory / TRACKING_WORKBOOK_NAME
    source_path = json_path
    final_path: Path | None = None

    try:
        if not json_path.exists():
            raise FileNotFoundError(
                f"找不到结构化简历JSON：{json_path.resolve()}"
            )

        if json_path.suffix.lower() != ".json":
            raise ValueError("input_file必须指向JSON文件。")

        if not TEMPLATE_PATH.exists():
            raise FileNotFoundError(
                f"找不到标准简历模板：{TEMPLATE_PATH.resolve()}"
            )

        data = json.loads(json_path.read_text(encoding="utf-8-sig"))

        if not isinstance(data, dict):
            raise ValueError("结构化简历JSON顶层必须是对象。")

        source_value = data.get("source_file", "")

        if source_value:
            source_path = Path(source_value)

        output_name = _output_stem(data, json_path)
        final_directory = output_root / FINAL_RESUMES_DIRNAME
        final_directory.mkdir(parents=True, exist_ok=True)
        final_path = final_directory / f"{output_name}_标准简历.docx"

        _generate_document(
            data=data,
            output_path=final_path,
            with_photo=bool(with_photo),
            remove_headers=remove_headers,
        )
        update_supplement_report(
            data=data,
            final_docx_path=final_path,
            workbook_path=final_directory / SUPPLEMENT_WORKBOOK_NAME,
        )

        if tracking_workbook.exists():
            update_generation_record(
                workbook_path=tracking_workbook,
                source_path=source_path,
                json_path=json_path,
                docx_path=final_path,
                status="成功",
            )

        print(f"  [Word生成完成] {final_path.resolve()}")
        return str(final_path.resolve())

    except Exception as error:
        error_file = write_error_log(
            output_dir=process_directory,
            input_file_path=source_path,
            stage="标准简历生成",
        )

        if tracking_workbook.exists():
            try:
                update_generation_record(
                    workbook_path=tracking_workbook,
                    source_path=source_path,
                    json_path=json_path,
                    docx_path=(
                        final_path
                        if final_path and final_path.exists()
                        else None
                    ),
                    status="失败",
                    error_message=str(error),
                )
            except Exception:
                # 过程记录失败不能掩盖最初的Word生成异常。
                pass

        print(f"  [Word生成失败] {error}")
        print(f"  详细错误已写入：{error_file.resolve()}")
        return None


def _generate_document(
    data: dict[str, Any],
    output_path: Path,
    with_photo: bool,
    remove_headers: bool = False,
) -> None:
    """复制模板并按固定分区顺序写入结构化数据。"""

    document = Document(TEMPLATE_PATH)
    section_elements = _section_elements(document)
    _clear_document_body(document)
    if remove_headers:
        _remove_headers(document)
    _set_document_defaults(document)

    _append_section(document, section_elements["基本资料"])
    _add_basic_information(
        document,
        _mapping(data.get("basic_information")),
        with_photo,
    )

    _append_section(document, section_elements["专业技能"])
    _add_skills(document, _text(data.get("professional_skills")))

    _append_section(document, section_elements["工作经历"])
    _add_work_experiences(
        document,
        _list(data.get("work_experiences")),
    )

    project_section = deepcopy(section_elements["项目经验"])
    _replace_xml_text(
        project_section,
        "Working Experience",
        "Project Experience",
    )
    _append_section(document, project_section)
    _add_project_experiences(
        document,
        _list(data.get("project_experiences")),
    )

    _append_section(document, section_elements["教育经历"])
    _add_education_experiences(
        document,
        _list(data.get("education_experiences")),
    )

    _append_section(document, section_elements["工作年限"])
    _add_work_years(document, _mapping(data.get("work_years")))

    _highlight_placeholders(document)
    _warn_markdown_residues(document)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)
    _force_microsoft_yahei(output_path)


def _section_elements(document: Document) -> dict[str, Any]:
    """从清洁模板中定位并复制六个带图标的分区标题。"""

    result: dict[str, Any] = {}

    for paragraph in document.paragraphs:
        text = _xml_text(paragraph._p)

        for label in SECTION_LABELS:
            if label in text and label not in result:
                result[label] = deepcopy(paragraph._p)

    missing = [label for label in SECTION_LABELS if label not in result]

    if missing:
        raise RuntimeError(
            "标准模板缺少分区组件：" + "、".join(missing)
        )

    return result


def _clear_document_body(document: Document) -> None:
    """清除模板正文，仅保留节属性、页眉和页脚关系。"""

    body = document._element.body

    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def _remove_headers(document: Document) -> None:
    """清空页眉部件并解除所有节的页眉引用。"""

    header_parts: dict[str, Any] = {}
    for relationship in document.part.rels.values():
        if relationship.reltype == RELATIONSHIP_TYPE.HEADER:
            header_parts[str(relationship.target_part.partname)] = (
                relationship.target_part
            )

    for header_part in header_parts.values():
        header_element = header_part.element
        for child in list(header_element):
            header_element.remove(child)
        # 保留一个合法空段落，避免部分Word版本修复空页眉部件。
        header_element.append(OxmlElement("w:p"))
        for relationship_id in list(header_part.rels):
            header_part.drop_rel(relationship_id)

    for section in document.sections:
        section_properties = section._sectPr
        for reference in list(
            section_properties.findall(qn("w:headerReference"))
        ):
            section_properties.remove(reference)


def _append_section(document: Document, element: Any) -> None:
    """插入模板原有的图标化分区标题。"""

    body = document._element.body
    sect_pr = body.sectPr
    section_element = deepcopy(element)
    paragraph_properties = section_element.get_or_add_pPr()

    if paragraph_properties.find(qn("w:keepNext")) is None:
        paragraph_properties.append(OxmlElement("w:keepNext"))

    body.insert(body.index(sect_pr), section_element)


def _set_document_defaults(document: Document) -> None:
    """设置A4版心，以及正文默认的微软雅黑和单倍行距。"""

    for section in document.sections:
        section.page_width = Cm(21)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(2.54)
        section.bottom_margin = Cm(2.54)
        section.left_margin = Cm(1.91)
        section.right_margin = Cm(1.91)

    for style in document.styles:
        if hasattr(style, "font"):
            style.font.name = FONT_NAME
            _set_style_east_asia_font(style)

    document.styles["Normal"].paragraph_format.line_spacing = (
        BODY_LINE_SPACING
    )


def _add_basic_information(
    document: Document,
    basic: dict[str, Any],
    with_photo: bool,
) -> None:
    """生成三行两列或带纵向照片占位的三行三列表格。"""

    columns = 3 if with_photo else 2
    table = document.add_table(rows=3, cols=columns)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    _remove_table_borders(table)

    width_values = (
        (6.872, 6.872, 3.436)
        if with_photo
        else (8.59, 8.59)
    )
    _set_table_grid(table, width_values)

    for row in table.rows:
        for cell, width in zip(row.cells, width_values):
            cell.width = Cm(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            _set_cell_margins(
                cell,
                top=80,
                start=180,
                bottom=80,
                end=180,
            )

    values = (
        (
            ("姓名", _text(basic.get("name"))),
            ("性别", _text(basic.get("gender"))),
        ),
        (
            ("出生年份", _text(basic.get("birth_year"))),
            ("籍贯", _text(basic.get("native_place"))),
        ),
        (("岗位职称", PLACEHOLDER), None),
    )

    for row_index, row_values in enumerate(values):
        for column_index, field in enumerate(row_values):
            _fill_basic_cell(
                table.cell(row_index, column_index),
                field,
                first_row=row_index == 0,
            )

    if with_photo:
        photo_cells = [table.cell(row_index, 2) for row_index in range(3)]
        _set_cell_border(
            photo_cells[0],
            color="B7B7B7",
            size="8",
            edges=("top", "left", "right"),
        )
        _set_cell_border(
            photo_cells[1],
            color="B7B7B7",
            size="8",
            edges=("left", "right"),
        )
        _set_cell_border(
            photo_cells[2],
            color="B7B7B7",
            size="8",
            edges=("bottom", "left", "right"),
        )
        photo_cell = photo_cells[0].merge(photo_cells[2])
        paragraph = photo_cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _set_body_line_spacing(paragraph)
        run = paragraph.add_run("证件照\n占位")
        _format_run(run, FIVE, BODY_COLOR, bold=False)


def _fill_basic_cell(
    cell,
    field: tuple[str, str] | None,
    first_row: bool,
) -> None:
    """写入基本资料的单个字段；空字段保留为空白单元格。"""

    cell.text = ""
    paragraph = cell.paragraphs[0]
    _set_body_line_spacing(paragraph)
    paragraph.paragraph_format.space_after = Pt(0)

    if first_row:
        paragraph.paragraph_format.space_before = SECTION_CONTENT_SPACE_BEFORE

    if field is None:
        return

    label, value = field
    run = paragraph.add_run(f"{label}：{value or PLACEHOLDER}")
    _format_run(run, FIVE, BODY_COLOR, bold=False)


def _add_skills(document: Document, skills: str) -> None:
    """按通用Markdown语义写入专业技能。"""

    _add_markdown_content(
        document=document,
        text=skills or PLACEHOLDER,
        size=FIVE,
        color=BODY_COLOR,
        first_space_before=SECTION_CONTENT_SPACE_BEFORE,
    )


def _add_work_experiences(
    document: Document,
    experiences: list[Any],
) -> None:
    """用统一列宽排列时间、公司和岗位，保证各列起点一致。"""

    rows = [_mapping(item) for item in experiences]

    if not rows:
        rows = [
            {
                "time": "时间待补充",
                "company_name": "公司名称待补充",
                "position_name": "岗位待补充",
                "description": "",
            }
        ]

    rows = [
        {
            **item,
            "time": _text(item.get("time")) or "时间待补充",
            "company_name": (
                _text(item.get("company_name"))
                or "公司名称待补充"
            ),
            "position_name": (
                _text(item.get("position_name"))
                or "岗位待补充"
            ),
            "description": _text(item.get("description")),
        }
        for item in rows
    ]

    time_width, position_width, company_width = _work_column_widths(rows)

    for index, item in enumerate(rows):
        if index:
            _add_work_table_separator(document)

        # 每段工作使用独立表格。若全部经历共用一个表格，Word会把
        # 不同标题行的keepNext串成分页链，造成整段误移或标题孤行。
        table = document.add_table(rows=0, cols=3)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        _remove_table_borders(table)
        _set_table_grid(
            table,
            (time_width, company_width, position_width),
        )
        title_row = table.add_row()
        title_row._tr.get_or_add_trPr().append(_cant_split_element())
        values = (
            _text(item.get("time")),
            _text(item.get("company_name")),
            _text(item.get("position_name")),
        )

        for cell, value, width in zip(
            title_row.cells,
            values,
            (time_width, company_width, position_width),
        ):
            cell.width = Cm(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            _set_cell_margins(cell, top=0, start=0, bottom=0, end=0)
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            _set_body_line_spacing(paragraph)
            paragraph.paragraph_format.space_after = Pt(0)
            if index == 0:
                paragraph.paragraph_format.space_before = (
                    SECTION_CONTENT_SPACE_BEFORE
                )
            elif (
                _text(rows[index - 1].get("description"))
                and _text(item.get("description"))
            ):
                # 使用段前距形成经历间的视觉空行。不要插入空表格行，
                # 否则它与“与下段同页”组合时会触发Word异常整块分页。
                paragraph.paragraph_format.space_before = (
                    WORK_ENTRY_SPACE_BEFORE
                )
            paragraph.paragraph_format.keep_with_next = (
                bool(_text(item.get("description")))
            )
            run = paragraph.add_run(value)
            _format_run(run, SMALL_FOUR, ACCENT_COLOR, bold=True)

        description = _text(item.get("description"))

        if description:
            description_cell = table.add_row().cells[0]
            description_cell = description_cell.merge(
                table.rows[-1].cells[2]
            )
            _set_cell_margins(
                description_cell,
                top=70,
                start=0,
                bottom=0,
                end=0,
            )
            label_paragraph = description_cell.paragraphs[0]
            _set_body_line_spacing(label_paragraph)
            label_paragraph.paragraph_format.space_after = Pt(0)
            label_paragraph.paragraph_format.keep_with_next = True
            label_run = label_paragraph.add_run("工作内容：")
            _format_run(label_run, FIVE, BODY_COLOR, bold=True)

            # 标签单独占一行，Markdown描述从下一段开始；避免控制符
            # 影响“工作内容：”的固定格式。
            _add_markdown_content(
                cell=description_cell,
                text=description,
                size=FIVE,
                color=BODY_COLOR,
                append=True,
            )


def _add_work_table_separator(document: Document) -> None:
    """分隔相邻工作表格，避免Word自动合并且不产生可见空行。"""

    paragraph = document.add_paragraph()
    _set_body_line_spacing(paragraph)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = Pt(1)


def _work_column_widths(
    experiences: list[dict[str, Any]],
) -> tuple[float, float, float]:
    """按最长文本估算时间与岗位列宽，时间列额外留两个中文空格。"""

    total_width = 17.18
    time_units = max(
        (_display_units(_text(item.get("time"))) for item in experiences),
        default=0,
    ) + 4
    position_units = max(
        (
            _display_units(_text(item.get("position_name")))
            for item in experiences
        ),
        default=0,
    )
    time_width = min(max(time_units * 0.21, 2.8), 5.2)
    # 微软雅黑中文实际字宽略大于简单字符估算，额外留出安全余量。
    position_width = min(max(position_units * 0.235 + 0.25, 2.4), 5.0)
    company_width = total_width - time_width - position_width

    if company_width < 4.5:
        deficit = 4.5 - company_width
        time_width = max(2.8, time_width - deficit / 2)
        position_width = max(2.4, position_width - deficit / 2)
        company_width = total_width - time_width - position_width

    return time_width, position_width, company_width


def _add_project_experiences(
    document: Document,
    experiences: list[Any],
) -> None:
    """生成带中文序号的项目标题和必需的项目描述区域。"""

    rows = [_mapping(item) for item in experiences]

    if not rows:
        rows = [
            {
                "project_name": "项目名称待补充",
                "position_name": "",
                "time": "",
                "description": "项目描述：待补充",
            }
        ]

    for index, item in enumerate(rows, start=1):
        title_parts = [
            f"项目{_chinese_number(index)}："
            f"{_text(item.get('project_name')) or '项目名称待补充'}"
        ]

        for optional_value in (
            _text(item.get("position_name")),
            _text(item.get("time")),
        ):
            if optional_value:
                title_parts.append(optional_value)

        title = " | ".join(title_parts)
        title_paragraph = document.add_paragraph()
        _set_body_line_spacing(title_paragraph)
        title_paragraph.paragraph_format.keep_with_next = True
        if index == 1:
            title_paragraph.paragraph_format.space_before = (
                PROJECT_CONTENT_SPACE_BEFORE
            )
        title_paragraph.paragraph_format.space_after = Pt(0)
        run = title_paragraph.add_run(title)
        _format_run(run, SMALL_FOUR, ACCENT_COLOR, bold=True)

        description = (
            _text(item.get("description"))
            or "项目描述：待补充"
        )
        _add_markdown_content(
            document=document,
            text=description,
            size=FIVE,
            color=BODY_COLOR,
        )

        if index < len(rows):
            document.add_paragraph()


def _add_education_experiences(
    document: Document,
    experiences: list[Any],
) -> None:
    """按时间、学校、学历、专业顺序生成教育经历。"""

    rows = [_mapping(item) for item in experiences]

    if not rows:
        rows = [
            {
                "time": "时间待补充",
                "school_name": "学校待补充",
                "degree": "学历待补充",
                "major": "专业待补充",
            }
        ]

    for index, item in enumerate(rows):
        values = [
            _text(item.get("time")) or "时间待补充",
            _text(item.get("school_name")) or "学校待补充",
            _text(item.get("degree")) or "学历待补充",
            _text(item.get("major")) or "专业待补充",
        ]
        paragraph = document.add_paragraph()
        _set_body_line_spacing(paragraph)
        if index == 0:
            paragraph.paragraph_format.space_before = (
                SECTION_CONTENT_SPACE_BEFORE
            )
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(
            " | ".join(value for value in values if value)
        )
        _format_run(run, SMALL_FOUR, ACCENT_COLOR, bold=True)


def _add_work_years(document: Document, work_years: dict[str, Any]) -> None:
    """写入向下取整后的工作年限，不自行重新计算。"""

    value = work_years.get("value")
    display_value = ""

    if isinstance(value, (int, float)) and not isinstance(value, bool):
        number = float(value)

        if math.isfinite(number):
            display_value = (
                str(int(number)) if number.is_integer() else str(number)
            ) + "年"

    paragraph = document.add_paragraph()
    _set_body_line_spacing(paragraph)
    paragraph.paragraph_format.space_before = SECTION_CONTENT_SPACE_BEFORE
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(display_value or PLACEHOLDER)
    _format_run(run, SMALL_FOUR, ACCENT_COLOR, bold=True)


def _fill_markdown_cell(cell, text: str, size, color) -> None:
    """在合并单元格中按通用Markdown语义写入描述。"""

    cell.text = ""
    _add_markdown_content(
        cell=cell,
        text=text,
        size=size,
        color=color,
    )


def _add_markdown_content(
    text: str,
    size,
    color,
    document: Document | None = None,
    cell=None,
    first_space_before=None,
    append: bool = False,
) -> None:
    """把Markdown语义块统一写入正文或表格单元格。"""

    if (document is None) == (cell is None):
        raise ValueError("document和cell必须且只能提供一个。")

    blocks = parse_resume_markdown(text)

    if not blocks:
        blocks = parse_resume_markdown(PLACEHOLDER)

    list_number_ids: dict[tuple[str, int], int] = {}

    for index, block in enumerate(blocks):
        if cell is not None:
            paragraph = (
                cell.paragraphs[0]
                if index == 0 and not append
                else cell.add_paragraph()
            )
        else:
            paragraph = document.add_paragraph()

        _format_markdown_paragraph(
            paragraph,
            block,
            size,
            color,
            list_number_ids,
            first_space_before if index == 0 else None,
        )


def _format_markdown_paragraph(
    paragraph,
    block: MarkdownBlock,
    size,
    color,
    list_number_ids: dict[tuple[str, int], int],
    first_space_before=None,
) -> None:
    """设置Markdown块的段落语义并写入全部内联格式。"""

    _set_body_line_spacing(paragraph)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.space_before = (
        first_space_before
        if first_space_before is not None
        else Pt(4)
        if block.kind == "heading"
        else Pt(0)
    )

    if block.kind == "heading":
        paragraph.paragraph_format.keep_with_next = True
    elif block.kind == "quote":
        paragraph.paragraph_format.left_indent = Cm(
            0.45 * max(1, block.quote_depth)
        )
    elif block.kind == "code":
        paragraph.paragraph_format.left_indent = Cm(0.45)
    elif block.kind in {"bullet_list", "ordered_list"}:
        _apply_list_numbering(paragraph, block, list_number_ids)

    for fragment in block.fragments:
        if fragment.link:
            run = _add_hyperlink_run(
                paragraph,
                fragment.text,
                fragment.link,
            )
        else:
            run = paragraph.add_run()
            _write_run_text(run, fragment.text)

        _format_run(
            run,
            size,
            color,
            bold=fragment.bold or block.kind == "heading",
        )
        run.font.italic = fragment.italic
        run.font.strike = fragment.strike
        if fragment.link:
            run.font.underline = True


def _write_run_text(run, text: str) -> None:
    """把Markdown软换行转换为Word行内换行，而不是额外空段落。"""

    parts = text.split("\n")
    for index, part in enumerate(parts):
        if index:
            run.add_break()
        if part:
            run.add_text(part)


def _add_hyperlink_run(paragraph, text: str, url: str):
    """创建保留可见文字的外部Word超链接。"""

    relationship_id = paragraph.part.relate_to(
        url,
        RELATIONSHIP_TYPE.HYPERLINK,
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run_element = OxmlElement("w:r")
    hyperlink.append(run_element)
    paragraph._p.append(hyperlink)
    run = Run(run_element, paragraph)
    _write_run_text(run, text)
    return run


def _apply_list_numbering(
    paragraph,
    block: MarkdownBlock,
    list_number_ids: dict[tuple[str, int], int],
) -> None:
    """为无序和有序列表创建真实Word编号，而不是显示Markdown标记。"""

    key = (block.kind, block.list_id)
    num_id = list_number_ids.get(key)

    if num_id is None:
        num_id = _create_numbering_definition(
            paragraph,
            ordered=block.kind == "ordered_list",
            start=block.list_start,
        )
        list_number_ids[key] = num_id

    paragraph_properties = paragraph._p.get_or_add_pPr()
    number_properties = paragraph_properties.find(qn("w:numPr"))
    if number_properties is None:
        number_properties = OxmlElement("w:numPr")
        paragraph_properties.append(number_properties)

    level_element = OxmlElement("w:ilvl")
    level_element.set(qn("w:val"), str(min(block.level, 8)))
    number_id_element = OxmlElement("w:numId")
    number_id_element.set(qn("w:val"), str(num_id))
    number_properties.append(level_element)
    number_properties.append(number_id_element)


def _create_numbering_definition(
    paragraph,
    ordered: bool,
    start: int,
) -> int:
    """向当前文档加入一套九级列表编号定义。"""

    numbering = _ensure_numbering_part(paragraph.part).element
    abstract_ids = [
        int(element.get(qn("w:abstractNumId")))
        for element in numbering.findall(qn("w:abstractNum"))
    ]
    number_ids = [
        int(element.get(qn("w:numId")))
        for element in numbering.findall(qn("w:num"))
    ]
    abstract_id = max(abstract_ids, default=-1) + 1
    number_id = max(number_ids, default=0) + 1
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi_level = OxmlElement("w:multiLevelType")
    multi_level.set(qn("w:val"), "multilevel")
    abstract.append(multi_level)

    for level in range(9):
        level_element = OxmlElement("w:lvl")
        level_element.set(qn("w:ilvl"), str(level))
        start_element = OxmlElement("w:start")
        start_element.set(
            qn("w:val"),
            str(start if ordered and level == 0 else 1),
        )
        number_format = OxmlElement("w:numFmt")
        number_format.set(
            qn("w:val"),
            "decimal" if ordered else "bullet",
        )
        level_text = OxmlElement("w:lvlText")
        level_text.set(
            qn("w:val"),
            f"%{level + 1}." if ordered else "·",
        )
        justification = OxmlElement("w:lvlJc")
        justification.set(qn("w:val"), "left")
        paragraph_properties = OxmlElement("w:pPr")
        indentation = OxmlElement("w:ind")
        indentation.set(qn("w:left"), str(420 + level * 360))
        indentation.set(qn("w:hanging"), "240")
        paragraph_properties.append(indentation)
        run_properties = OxmlElement("w:rPr")
        fonts = OxmlElement("w:rFonts")
        for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
            fonts.set(qn(f"w:{attribute}"), FONT_NAME)
        run_properties.append(fonts)

        for child in (
            start_element,
            number_format,
            level_text,
            justification,
            paragraph_properties,
            run_properties,
        ):
            level_element.append(child)
        abstract.append(level_element)

    numbering.insert(0, abstract)
    number = OxmlElement("w:num")
    number.set(qn("w:numId"), str(number_id))
    abstract_reference = OxmlElement("w:abstractNumId")
    abstract_reference.set(qn("w:val"), str(abstract_id))
    number.append(abstract_reference)
    numbering.append(number)
    return number_id


def _ensure_numbering_part(document_part) -> NumberingPart:
    """为未携带numbering.xml的模板创建标准Word编号部件。"""

    try:
        return document_part.part_related_by(
            RELATIONSHIP_TYPE.NUMBERING
        )
    except KeyError:
        numbering_part = NumberingPart(
            PackURI("/word/numbering.xml"),
            CONTENT_TYPE.WML_NUMBERING,
            OxmlElement("w:numbering"),
            document_part.package,
        )
        document_part.relate_to(
            numbering_part,
            RELATIONSHIP_TYPE.NUMBERING,
        )
        return numbering_part


def _format_run(run, size, color, bold: bool) -> None:
    """应用正文统一字体与角色格式。"""

    run.font.name = FONT_NAME
    run.font.size = size
    run.font.bold = bold
    run.font.color.rgb = color
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.get_or_add_rFonts()

    for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
        r_fonts.set(qn(f"w:{attribute}"), FONT_NAME)


def _highlight_placeholders(document: Document) -> None:
    """只把最终Word中占位词本身设为红色，并保留周围格式。"""

    # 占位词可能嵌在“时间待补充”等完整字段中，因此需要拆分Word游程，
    # 不能把整行文字统一设为红色。
    for run_element in list(document._element.body.iter(qn("w:r"))):
        run = Run(run_element, None)
        text = run.text

        if PLACEHOLDER not in text:
            continue

        parent = run_element.getparent()
        insert_at = parent.index(run_element)
        segments = text.split(PLACEHOLDER)

        for index, segment in enumerate(segments):
            if segment:
                clone = deepcopy(run_element)
                clone_run = Run(clone, None)
                clone_run.text = segment
                parent.insert(insert_at, clone)
                insert_at += 1

            if index < len(segments) - 1:
                clone = deepcopy(run_element)
                clone_run = Run(clone, None)
                clone_run.text = PLACEHOLDER
                clone_run.font.color.rgb = PLACEHOLDER_COLOR
                parent.insert(insert_at, clone)
                insert_at += 1

        parent.remove(run_element)


def _set_style_east_asia_font(style) -> None:
    """补齐样式中的中西文字体声明。"""

    try:
        r_pr = style.element.get_or_add_rPr()
        r_fonts = r_pr.get_or_add_rFonts()
    except AttributeError:
        return

    for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
        r_fonts.set(qn(f"w:{attribute}"), FONT_NAME)


def _remove_table_borders(table) -> None:
    """把布局表设置为无边框。"""

    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")

    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)

    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))

        if element is None:
            element = OxmlElement(tag)
            borders.append(element)

        element.set(qn("w:val"), "nil")


def _set_body_line_spacing(paragraph) -> None:
    """设置真实单倍行距，并关闭模板文档网格对正文的拉伸。"""

    paragraph.paragraph_format.line_spacing = BODY_LINE_SPACING
    paragraph_properties = paragraph._p.get_or_add_pPr()
    snap_to_grid = paragraph_properties.find(qn("w:snapToGrid"))

    if snap_to_grid is None:
        snap_to_grid = OxmlElement("w:snapToGrid")
        paragraph_properties.append(snap_to_grid)

    snap_to_grid.set(qn("w:val"), "0")


def _set_cell_border(
    cell,
    color: str,
    size: str,
    edges: tuple[str, ...],
) -> None:
    """为纵向照片占位区域的指定边设置细边框。"""

    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")

    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)

    for edge in edges:
        existing = borders.find(qn(f"w:{edge}"))

        if existing is not None:
            borders.remove(existing)

        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)
        borders.append(element)


def _set_cell_margins(
    cell,
    top: int,
    start: int,
    bottom: int,
    end: int,
) -> None:
    """以DXA设置单元格内边距。"""

    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")

    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)

    for edge, value in (
        ("top", top),
        ("start", start),
        ("bottom", bottom),
        ("end", end),
    ):
        element = tc_mar.find(qn(f"w:{edge}"))

        if element is None:
            element = OxmlElement(f"w:{edge}")
            tc_mar.append(element)

        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def _set_table_grid(table, widths_cm: tuple[float, ...]) -> None:
    """同步表格网格和单元格宽度，避免Word自动重新分栏。"""

    table_grid = table._tbl.tblGrid

    for child in list(table_grid):
        table_grid.remove(child)

    for width in widths_cm:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(int(Cm(width).twips)))
        table_grid.append(grid_col)


def _cant_split_element():
    """禁止工作经历标题行跨页拆分。"""

    return OxmlElement("w:cantSplit")


def _display_units(text: str) -> int:
    """以半角字符为1、全角字符为2估算显示宽度。"""

    return sum(
        2
        if unicodedata.east_asian_width(character) in {"W", "F", "A"}
        else 1
        for character in text
    )


def _chinese_number(number: int) -> str:
    """把正整数转换为项目标题使用的中文序号。"""

    digits = "零一二三四五六七八九"

    if number < 10:
        return digits[number]

    if number < 20:
        return "十" + (digits[number % 10] if number % 10 else "")

    if number < 100:
        tens, ones = divmod(number, 10)
        return digits[tens] + "十" + (digits[ones] if ones else "")

    return str(number)


def _output_stem(data: dict[str, Any], json_path: Path) -> str:
    """优先使用原始简历stem，缺失时回退到JSON文件名。"""

    source_file = _text(data.get("source_file"))

    if source_file:
        return Path(source_file).stem

    stem = json_path.stem
    return stem.removesuffix("_extracted") or "标准简历"


def _xml_text(element: Any) -> str:
    """读取普通段落和文本框中的全部Word文本。"""

    return "".join(
        node.text or "" for node in element.iter(qn("w:t"))
    )


def _replace_xml_text(element: Any, old: str, new: str) -> None:
    """同时修正DrawingML和VML回退文本中的英文标题。"""

    for node in element.iter(qn("w:t")):
        if node.text and old in node.text:
            node.text = node.text.replace(old, new)


def _force_microsoft_yahei(docx_path: Path) -> None:
    """在OOXML层强制所有WordprocessingML文本使用微软雅黑。"""

    _force_word_font(docx_path, FONT_NAME)


def _force_word_font(docx_path: Path, font_name: str) -> None:
    """在OOXML层强制文档全部文本使用指定字体。"""

    patched_path = docx_path.with_name(
        f".{docx_path.stem}.font-patched.docx"
    )

    try:
        with ZipFile(docx_path, "r") as source, ZipFile(
            patched_path,
            "w",
            compression=ZIP_DEFLATED,
        ) as target:
            for item in source.infolist():
                content = source.read(item.filename)

                if (
                    item.filename.startswith("word/")
                    and item.filename.endswith(".xml")
                ):
                    content = _patch_xml_fonts(content, font_name)

                target.writestr(item, content)

        shutil.move(str(patched_path), str(docx_path))
    finally:
        if patched_path.exists():
            patched_path.unlink()


def _patch_xml_fonts(
    content: bytes,
    font_name: str = FONT_NAME,
) -> bytes:
    """只修改字体属性，不重排其他XML结构。"""

    from lxml import etree

    parser = etree.XMLParser(remove_blank_text=False)
    root = etree.fromstring(content, parser)
    namespaces = {
        "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
        "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    }

    for fonts in root.xpath("//w:rFonts", namespaces=namespaces):
        for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
            fonts.set(qn(f"w:{attribute}"), font_name)

        for attribute in (
            "asciiTheme",
            "hAnsiTheme",
            "eastAsiaTheme",
            "cstheme",
        ):
            fonts.attrib.pop(qn(f"w:{attribute}"), None)

    for run_properties in root.xpath(
        "//a:rPr | //a:defRPr | //a:endParaRPr",
        namespaces=namespaces,
    ):
        for tag in ("latin", "ea", "cs"):
            child = run_properties.find(f"{{{namespaces['a']}}}{tag}")

            if child is None:
                child = etree.SubElement(
                    run_properties,
                    f"{{{namespaces['a']}}}{tag}",
                )

            child.set("typeface", font_name)

    return etree.tostring(
        root,
        xml_declaration=True,
        encoding="UTF-8",
        standalone=True,
    )


def _warn_markdown_residues(document: Document) -> None:
    """生成前扫描可见文本；警告只提示，不中断Word输出。"""

    texts = [
        _xml_text(paragraph)
        for paragraph in document._element.body.iter(qn("w:p"))
    ]
    residues = find_markdown_residues(texts)

    if not residues:
        return

    print("  警告：最终Word中仍检测到可能的Markdown残留：")
    for text in residues[:10]:
        print(f"    - {text}")
    if len(residues) > 10:
        print(f"    - 其余 {len(residues) - 10} 处已省略")


def _mapping(value: Any) -> dict[str, Any]:
    return value if isinstance(value, dict) else {}


def _list(value: Any) -> list[Any]:
    return value if isinstance(value, list) else []


def _text(value: Any) -> str:
    return value.strip() if isinstance(value, str) else ""
