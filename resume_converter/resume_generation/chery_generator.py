"""根据结构化JSON生成奇瑞宋体标准简历。"""

from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from utils.error_logging import write_error_log
from utils.processing_records import (
    TRACKING_WORKBOOK_NAME,
    update_generation_record,
)

from .generator import (
    FINAL_RESUMES_DIRNAME,
    PROCESS_DATA_DIRNAME,
    _add_markdown_content,
    _clear_document_body,
    _display_units,
    _force_word_font,
    _highlight_placeholders,
    _output_stem,
    _set_body_line_spacing,
    _warn_markdown_residues,
)
from .supplement_report import (
    SUPPLEMENT_WORKBOOK_NAME,
    collect_chery_supplement_items,
    update_supplement_report,
)


FONT_NAME = "宋体"
HEADING_SIZE = Pt(14)
BODY_SIZE = Pt(12)
NOTE_SIZE = Pt(10.5)
BODY_COLOR = RGBColor(0, 0, 0)
PLACEHOLDER_COLOR = RGBColor(255, 0, 0)
NOTE_COLOR = RGBColor(192, 0, 0)
PLACEHOLDER = "待补充"
AI_REFERENCE_NOTE = "【AI根据简历已有信息生成，仅供参考，请人工确认】"
WORK_COLUMN_GAP = 3


def CheryResumeGenerator(
    input_file: str,
    output_dir: str,
) -> str | None:
    """将奇瑞结构化提取JSON生成宋体标准简历。"""

    json_path = Path(input_file)
    output_root = Path(output_dir)
    process_directory = output_root / PROCESS_DATA_DIRNAME
    tracking_workbook = process_directory / TRACKING_WORKBOOK_NAME
    source_path = json_path
    final_path: Path | None = None

    try:
        if not json_path.exists():
            raise FileNotFoundError(
                f"找不到结构化简历JSON：{json_path.resolve()}"
            )
        if json_path.suffix.lower() != ".json":
            raise ValueError("input_file必须指向JSON文件。")

        data = json.loads(json_path.read_text(encoding="utf-8-sig"))
        if not isinstance(data, dict):
            raise ValueError("结构化简历JSON顶层必须是对象。")

        source_value = _text(data.get("source_file"))
        if source_value:
            source_path = Path(source_value)

        final_directory = output_root / FINAL_RESUMES_DIRNAME
        final_directory.mkdir(parents=True, exist_ok=True)
        final_path = final_directory / (
            f"{_output_stem(data, json_path)}_标准简历.docx"
        )

        _generate_chery_document(data, final_path)
        update_supplement_report(
            data=data,
            final_docx_path=final_path,
            workbook_path=final_directory / SUPPLEMENT_WORKBOOK_NAME,
            supplement_items=collect_chery_supplement_items(data),
        )

        if tracking_workbook.exists():
            update_generation_record(
                workbook_path=tracking_workbook,
                source_path=source_path,
                json_path=json_path,
                docx_path=final_path,
                status="成功",
            )

        print(f"  [Word生成完成] {final_path.resolve()}")
        return str(final_path.resolve())

    except Exception as error:
        error_file = write_error_log(
            output_dir=process_directory,
            input_file_path=source_path,
            stage="奇瑞标准简历生成",
        )
        if tracking_workbook.exists():
            try:
                update_generation_record(
                    workbook_path=tracking_workbook,
                    source_path=source_path,
                    json_path=json_path,
                    docx_path=(
                        final_path
                        if final_path and final_path.exists()
                        else None
                    ),
                    status="失败",
                    error_message=str(error),
                )
            except Exception:
                pass

        print(f"  [Word生成失败] {error}")
        print(f"  详细错误已写入：{error_file.resolve()}")
        return None


def _generate_chery_document(
    data: dict[str, Any],
    output_path: Path,
) -> None:
    """按奇瑞示例的六分区顺序生成无装饰A4文档。"""

    document = Document()
    _clear_document_body(document)
    _set_document_defaults(document)

    _add_section_heading(document, "基本信息", first=True)
    _add_basic_information(document, data)

    _add_section_heading(document, "自我评价")
    _add_narrative(
        document,
        _text(data.get("self_evaluation")),
        _text(data.get("self_evaluation_source")),
    )

    _add_section_heading(document, "专业技能")
    _add_narrative(
        document,
        _text(data.get("professional_skills")),
        _text(data.get("professional_skills_source")),
    )

    _add_section_heading(document, "工作经历")
    _add_work_experiences(document, _list(data.get("work_experiences")))

    _add_section_heading(document, "项目经验")
    _add_project_experiences(
        document,
        _list(data.get("project_experiences")),
    )

    _add_section_heading(document, "教育背景")
    _add_education_experiences(
        document,
        _list(data.get("education_experiences")),
    )

    _highlight_placeholders(document)
    _warn_markdown_residues(document)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)
    _force_word_font(output_path, FONT_NAME)


def _set_document_defaults(document: Document) -> None:
    """复刻示例的A4版心，并设置宋体单倍行距。"""

    for section in document.sections:
        section.page_width = Cm(21)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(2.54)
        section.bottom_margin = Cm(2.54)
        section.left_margin = Cm(3.17)
        section.right_margin = Cm(3.17)
        section.header_distance = Cm(1.5)
        section.footer_distance = Cm(1.75)

    for style in document.styles:
        if not hasattr(style, "font"):
            continue
        style.font.name = FONT_NAME
        style_element = style.element
        run_properties = style_element.get_or_add_rPr()
        fonts = run_properties.get_or_add_rFonts()
        for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
            fonts.set(qn(f"w:{attribute}"), FONT_NAME)

    normal = document.styles["Normal"]
    normal.font.size = BODY_SIZE
    normal.paragraph_format.line_spacing = 1.0
    normal.paragraph_format.space_after = Pt(0)


def _add_section_heading(
    document: Document,
    text: str,
    first: bool = False,
) -> None:
    """添加四号加粗分区标题，并避免标题孤行。"""

    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(0 if first else 14)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.keep_with_next = True
    _set_body_line_spacing(paragraph)
    run = paragraph.add_run(text)
    _format_run(run, HEADING_SIZE, bold=True)


def _add_basic_information(
    document: Document,
    data: dict[str, Any],
) -> None:
    """逐行写入奇瑞基本信息及红色缺失占位。"""

    basic = _mapping(data.get("basic_information"))
    work_years = _mapping(data.get("work_years")).get("value")
    work_years_text = _format_years(work_years)

    for label, value in (
        ("姓名", _text(basic.get("name"))),
        ("性别", _text(basic.get("gender"))),
        ("工作年限", work_years_text),
        ("出生年月", _text(basic.get("birth_date"))),
        ("联系电话", _text(basic.get("phone"))),
        ("邮箱", _text(basic.get("email"))),
        ("籍贯", _text(basic.get("native_place"))),
    ):
        paragraph = document.add_paragraph()
        _set_body_line_spacing(paragraph)
        _add_labeled_value(paragraph, f"{label}：", value)


def _add_narrative(
    document: Document,
    text: str,
    source_type: str,
) -> None:
    """写入原文或AI参考长文本，缺失时保留红色占位。"""

    if not text:
        paragraph = document.add_paragraph()
        _set_body_line_spacing(paragraph)
        _add_placeholder(paragraph)
        return

    if source_type == "generated":
        note = document.add_paragraph()
        note.paragraph_format.space_after = Pt(3)
        _set_body_line_spacing(note)
        run = note.add_run(AI_REFERENCE_NOTE)
        _format_run(run, NOTE_SIZE, bold=True, color=NOTE_COLOR)

    _add_markdown_content(
        text=text,
        size=BODY_SIZE,
        color=BODY_COLOR,
        document=document,
    )


def _add_work_experiences(
    document: Document,
    experiences: list[Any],
) -> None:
    """以普通段落和空格补齐时间、公司、岗位的左对齐位置。"""

    normalized = [_mapping(item) for item in experiences]
    if not normalized:
        paragraph = document.add_paragraph()
        _set_body_line_spacing(paragraph)
        _add_placeholder(paragraph)
        return

    values = [
        (
            _dot_date(_text(item.get("time"))) or PLACEHOLDER,
            _text(item.get("company_name")) or PLACEHOLDER,
            _text(item.get("position_name")) or PLACEHOLDER,
        )
        for item in normalized
    ]
    max_time_units = max(_display_units(item[0]) for item in values)
    max_company_units = max(_display_units(item[1]) for item in values)

    for time_text, company_text, position_text in values:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.keep_together = True
        _set_body_line_spacing(paragraph)
        _format_run(paragraph.add_run(time_text), BODY_SIZE)
        _add_work_gap(paragraph, time_text, max_time_units)
        _format_run(paragraph.add_run(company_text), BODY_SIZE)
        _add_work_gap(paragraph, company_text, max_company_units)
        _format_run(paragraph.add_run(position_text), BODY_SIZE)


def _add_work_gap(paragraph, value: str, maximum_units: int) -> None:
    """补齐较短字段，并在最长字段后保留三个普通空格。"""

    padding = maximum_units - _display_units(value) + WORK_COLUMN_GAP
    _format_run(paragraph.add_run(" " * padding), BODY_SIZE)


def _add_project_experiences(
    document: Document,
    projects: list[Any],
) -> None:
    """按项目数字编号写入项目名称、描述和工作业绩。"""

    normalized = [_mapping(item) for item in projects]
    if not normalized:
        normalized = [{}]

    for index, project in enumerate(normalized, start=1):
        if index > 1:
            spacer = document.add_paragraph()
            spacer.paragraph_format.space_after = Pt(3)
            _set_body_line_spacing(spacer)

        title = document.add_paragraph()
        title.paragraph_format.keep_with_next = True
        _set_body_line_spacing(title)
        _format_run(title.add_run(f"项目{index}"), BODY_SIZE, bold=True)

        _add_project_name(
            document,
            _text(project.get("project_name")),
        )

        _add_project_block(
            document,
            "项目描述：",
            _text(project.get("description")),
        )
        _add_project_block(
            document,
            "工作业绩：",
            _text(project.get("achievement")),
        )


def _add_project_name(document: Document, text: str) -> None:
    """将项目名称作为独立字段写在项目编号的下一行。"""

    paragraph = document.add_paragraph()
    paragraph.paragraph_format.keep_with_next = True
    _set_body_line_spacing(paragraph)
    _format_run(paragraph.add_run("项目名称："), BODY_SIZE, bold=True)
    if text:
        _format_run(paragraph.add_run(text), BODY_SIZE)
    else:
        _add_placeholder(paragraph)


def _add_project_block(
    document: Document,
    label: str,
    text: str,
) -> None:
    """写入加粗项目子标题及对应Markdown正文。"""

    heading = document.add_paragraph()
    heading.paragraph_format.keep_with_next = True
    _set_body_line_spacing(heading)
    _format_run(heading.add_run(label), BODY_SIZE, bold=True)

    if text:
        _add_markdown_content(
            text=text,
            size=BODY_SIZE,
            color=BODY_COLOR,
            document=document,
        )
    else:
        _add_placeholder(heading)


def _add_education_experiences(
    document: Document,
    experiences: list[Any],
) -> None:
    """按时间、学校、专业、学历顺序写入教育背景。"""

    normalized = [_mapping(item) for item in experiences]
    if not normalized:
        normalized = [{}]

    for experience in normalized:
        paragraph = document.add_paragraph()
        _set_body_line_spacing(paragraph)
        values = (
            _dot_date(_text(experience.get("time"))),
            _text(experience.get("school_name")),
            _text(experience.get("major")),
            _text(experience.get("degree")),
        )
        for index, value in enumerate(values):
            if index:
                _format_run(paragraph.add_run(" | "), BODY_SIZE)
            if value:
                _format_run(paragraph.add_run(value), BODY_SIZE)
            else:
                _add_placeholder(paragraph)


def _add_labeled_value(paragraph, label: str, value: str) -> None:
    _format_run(paragraph.add_run(label), BODY_SIZE)
    if value:
        _format_run(paragraph.add_run(value), BODY_SIZE)
    else:
        _add_placeholder(paragraph)


def _add_placeholder(paragraph) -> None:
    _format_run(
        paragraph.add_run(PLACEHOLDER),
        BODY_SIZE,
        color=PLACEHOLDER_COLOR,
    )


def _format_run(
    run,
    size,
    bold: bool = False,
    color: RGBColor = BODY_COLOR,
) -> None:
    run.font.name = FONT_NAME
    run.font.size = size
    run.font.bold = bold
    run.font.color.rgb = color
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{attribute}"), FONT_NAME)


def _format_years(value: Any) -> str:
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        return ""
    numeric = float(value)
    return f"{int(numeric) if numeric.is_integer() else numeric:g}年"


def _dot_date(value: str) -> str:
    return value.replace("/", ".")


def _mapping(value: Any) -> dict[str, Any]:
    return value if isinstance(value, dict) else {}


def _list(value: Any) -> list[Any]:
    return value if isinstance(value, list) else []


def _text(value: Any) -> str:
    return value.strip() if isinstance(value, str) else ""
