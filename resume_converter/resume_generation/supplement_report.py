"""
维护最终标准简历对应的人工补充信息清单。

清单只记录会影响最终交付完整性的必填内容；工作描述、项目岗位和
项目时间等业务上允许缺失的字段不会被列入。
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo


SUPPLEMENT_WORKBOOK_NAME = "待补充信息.xlsx"
SUPPLEMENT_SHEET_NAME = "待补充信息"
SUPPLEMENT_HEADERS = (
    "候选人姓名",
    "最终简历路径",
    "需要补充的内容",
)

_BASIC_FIELDS = (
    ("name", "姓名"),
    ("gender", "性别"),
    ("birth_year", "出生年份"),
    ("native_place", "籍贯"),
)
_WORK_FIELDS = (
    ("time", "时间"),
    ("company_name", "公司名称"),
    ("position_name", "岗位名称"),
)
_PROJECT_FIELDS = (
    ("project_name", "项目名称"),
    ("description", "项目描述"),
)
_EDUCATION_FIELDS = (
    ("time", "就读时间"),
    ("school_name", "学校"),
    ("degree", "学历"),
    ("major", "专业"),
)


def update_supplement_report(
    data: dict[str, Any],
    final_docx_path: Path,
    workbook_path: Path,
) -> Path:
    """按最终Word绝对路径创建或更新一条人工补充记录。"""

    workbook_path.parent.mkdir(parents=True, exist_ok=True)
    workbook, worksheet = _open_or_create_workbook(workbook_path)
    final_path_value = str(final_docx_path.resolve())
    candidate_name = _text(
        _mapping(data.get("basic_information")).get("name")
    ) or "待补充"
    supplement_text = _format_supplement_items(
        collect_supplement_items(data)
    )
    target_row = None

    for row in range(2, worksheet.max_row + 1):
        if worksheet.cell(row, 2).value == final_path_value:
            target_row = row
            break

    if target_row is None:
        worksheet.append(
            (candidate_name, final_path_value, supplement_text)
        )
        target_row = worksheet.max_row
    else:
        worksheet.cell(target_row, 1, candidate_name)
        worksheet.cell(target_row, 3, supplement_text)

    _style_data_row(worksheet, target_row, supplement_text)
    _refresh_table(worksheet)
    workbook.save(workbook_path)
    workbook.close()
    return workbook_path.resolve()


def collect_supplement_items(data: dict[str, Any]) -> list[str]:
    """根据最终模板的必填规则汇总需要人工补充的内容。"""

    items: list[str] = []
    basic = _mapping(data.get("basic_information"))

    for field, label in _BASIC_FIELDS:
        if not _text(basic.get(field)):
            items.append(f"请补充{label}")

    # 岗位职称不由当前信息提取模块生成，始终交给人工确认。
    items.append("请确认岗位职称")

    if not _text(data.get("professional_skills")):
        items.append("请补充专业技能")

    work_experiences = [
        _mapping(item) for item in _list(data.get("work_experiences"))
    ]
    if not work_experiences:
        items.append("请补充工作经历：时间、公司名称、岗位名称")
    else:
        for index, experience in enumerate(work_experiences, start=1):
            missing = _missing_labels(experience, _WORK_FIELDS)
            if missing:
                items.append(
                    f"请补充第{index}段工作经历的{'、'.join(missing)}"
                )

    project_experiences = [
        _mapping(item)
        for item in _list(data.get("project_experiences"))
    ]
    if not project_experiences:
        items.append("请补充项目经历：项目名称、项目描述")
    else:
        for index, experience in enumerate(project_experiences, start=1):
            missing = _missing_labels(experience, _PROJECT_FIELDS)
            if missing:
                items.append(
                    f"请补充项目{_chinese_number(index)}的"
                    f"{'、'.join(missing)}"
                )

    education_experiences = [
        _mapping(item)
        for item in _list(data.get("education_experiences"))
    ]
    work_years_missing = _mapping(data.get("work_years")).get("value") is None

    if not education_experiences:
        message = "请补充教育经历：就读时间、学校、学历、专业"
        if work_years_missing:
            message += "；补充毕业时间后才能计算工作年限"
        items.append(message)
    else:
        missing_graduation_time = False
        for index, experience in enumerate(
            education_experiences,
            start=1,
        ):
            missing = _missing_labels(experience, _EDUCATION_FIELDS)
            if "就读时间" in missing:
                missing_graduation_time = True
            if missing:
                message = (
                    f"请补充第{index}段教育经历的{'、'.join(missing)}"
                )
                if work_years_missing and "就读时间" in missing:
                    message += "；补充毕业时间后才能计算工作年限"
                items.append(message)

        if work_years_missing and not missing_graduation_time:
            items.append("请确认最后一段学历的毕业时间，以便计算工作年限")

    return items


def _open_or_create_workbook(workbook_path: Path):
    """打开现有清单，或创建固定三列表格。"""

    if workbook_path.exists():
        workbook = load_workbook(workbook_path)
        if SUPPLEMENT_SHEET_NAME not in workbook.sheetnames:
            workbook.close()
            raise RuntimeError("待补充信息工作簿缺少同名工作表。")
        worksheet = workbook[SUPPLEMENT_SHEET_NAME]
        existing_headers = tuple(cell.value for cell in worksheet[1])
        if existing_headers != SUPPLEMENT_HEADERS:
            workbook.close()
            raise RuntimeError("待补充信息工作簿表头与当前版本不兼容。")
    else:
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = SUPPLEMENT_SHEET_NAME
        worksheet.append(SUPPLEMENT_HEADERS)

    _style_worksheet(worksheet)
    return workbook, worksheet


def _style_worksheet(worksheet) -> None:
    """设置便于招聘人员查看和筛选的固定样式。"""

    worksheet.freeze_panes = "A2"
    worksheet.sheet_view.showGridLines = False
    worksheet.row_dimensions[1].height = 26
    worksheet.column_dimensions["A"].width = 18
    worksheet.column_dimensions["B"].width = 62
    worksheet.column_dimensions["C"].width = 72
    header_fill = PatternFill(fill_type="solid", fgColor="1F4E78")

    for cell in worksheet[1]:
        cell.font = Font(name="微软雅黑", color="FFFFFF", bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

    for row in range(2, worksheet.max_row + 1):
        content = _text(worksheet.cell(row, 3).value)
        _style_data_row(worksheet, row, content)


def _style_data_row(worksheet, row: int, supplement_text: str) -> None:
    """设置数据行字体、换行与足够的显示高度。"""

    for cell in worksheet[row]:
        cell.font = Font(name="微软雅黑", size=10.5)
        cell.alignment = Alignment(vertical="top", wrap_text=True)

    line_count = max(1, supplement_text.count("\n") + 1)
    worksheet.row_dimensions[row].height = min(
        max(30, line_count * 20),
        300,
    )


def _refresh_table(worksheet) -> None:
    """让Excel结构化表始终覆盖全部记录。"""

    table_ref = f"A1:C{worksheet.max_row}"
    worksheet.auto_filter.ref = table_ref

    if "SupplementInformation" in worksheet.tables:
        worksheet.tables["SupplementInformation"].ref = table_ref
        return

    table = Table(displayName="SupplementInformation", ref=table_ref)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    worksheet.add_table(table)


def _format_supplement_items(items: list[str]) -> str:
    """把全部事项放入同一单元格，并保留清晰编号。"""

    if not items:
        return "无需补充"
    return "\n".join(
        f"{index}. {item}" for index, item in enumerate(items, start=1)
    )


def _missing_labels(
    value: dict[str, Any],
    fields: tuple[tuple[str, str], ...],
) -> list[str]:
    return [label for field, label in fields if not _text(value.get(field))]


def _chinese_number(number: int) -> str:
    """将常见项目序号转换为中文，超出范围时保留数字。"""

    digits = "零一二三四五六七八九"
    if number < 10:
        return digits[number]
    if number < 20:
        return "十" + (digits[number % 10] if number % 10 else "")
    if number < 100:
        tens, ones = divmod(number, 10)
        return digits[tens] + "十" + (digits[ones] if ones else "")
    return str(number)


def _mapping(value: Any) -> dict[str, Any]:
    return value if isinstance(value, dict) else {}


def _list(value: Any) -> list[Any]:
    return value if isinstance(value, list) else []


def _text(value: Any) -> str:
    return value.strip() if isinstance(value, str) else ""
