"""
处理记录工作簿的创建、追加和阶段状态更新工具。
"""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo


TRACKING_WORKBOOK_NAME = "processing_records.xlsx"
TRACKING_SHEET_NAME = "处理记录"
TRACKING_HEADERS = (
    "处理时间",
    "原始文件路径",
    "百度Markdown路径",
    "结构恢复Markdown路径",
    "文档解析状态",
    "文档解析错误",
    "结构化JSON路径",
    "信息提取状态",
    "信息提取错误",
    "标准简历Word路径",
    "Word生成状态",
    "Word生成错误",
)


def _style_worksheet(worksheet) -> None:
    """设置处理记录表的固定样式。"""

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = "A1:L1"
    worksheet.sheet_view.showGridLines = False
    worksheet.row_dimensions[1].height = 24

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="1F4E78",
    )

    for cell in worksheet[1]:
        cell.font = Font(
            color="FFFFFF",
            bold=True,
        )
        cell.fill = header_fill
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

    column_widths = {
        "A": 20,
        "B": 48,
        "C": 48,
        "D": 48,
        "E": 14,
        "F": 40,
        "G": 48,
        "H": 14,
        "I": 40,
        "J": 48,
        "K": 14,
        "L": 40,
    }

    for column, width in column_widths.items():
        worksheet.column_dimensions[column].width = width


def _open_or_create_workbook(workbook_path: Path):
    """打开现有处理记录，或创建兼容当前列结构的新工作簿。"""

    if workbook_path.exists():
        workbook = load_workbook(workbook_path)
        worksheet = workbook[TRACKING_SHEET_NAME]

        existing_headers = tuple(
            cell.value for cell in worksheet[1]
        )

        if existing_headers != TRACKING_HEADERS:
            # 兼容旧版六列表格，保留已有记录并扩展阶段字段。
            if existing_headers == (
                "处理时间",
                "原始文件路径",
                "百度Markdown路径",
                "结构恢复Markdown路径",
                "处理状态",
                "错误信息",
            ):
                worksheet.cell(1, 5, "文档解析状态")
                worksheet.cell(1, 6, "文档解析错误")
                worksheet.cell(1, 7, "结构化JSON路径")
                worksheet.cell(1, 8, "信息提取状态")
                worksheet.cell(1, 9, "信息提取错误")
                worksheet.cell(1, 10, "标准简历Word路径")
                worksheet.cell(1, 11, "Word生成状态")
                worksheet.cell(1, 12, "Word生成错误")
            elif existing_headers == TRACKING_HEADERS[:9]:
                # 兼容加入Word生成阶段之前的九列表格。
                worksheet.cell(1, 10, "标准简历Word路径")
                worksheet.cell(1, 11, "Word生成状态")
                worksheet.cell(1, 12, "Word生成错误")
            else:
                workbook.close()
                raise RuntimeError(
                    "处理记录表头与当前版本不兼容。"
                )
    else:
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = TRACKING_SHEET_NAME
        worksheet.append(TRACKING_HEADERS)

    _style_worksheet(worksheet)
    return workbook, worksheet


def _set_status_style(cell, status: str) -> None:
    """按阶段状态设置便于扫描的颜色。"""

    colors = {
        "成功": "E2F0D9",
        "失败": "FCE4D6",
        "待提取": "FFF2CC",
        "待生成": "FFF2CC",
    }
    cell.font = Font(bold=True)
    cell.alignment = Alignment(
        horizontal="center",
        vertical="top",
    )
    cell.fill = PatternFill(
        fill_type="solid",
        fgColor=colors.get(status, "FFFFFF"),
    )


def _refresh_table(worksheet) -> None:
    """让Excel结构化表覆盖全部现有记录。"""

    table_ref = f"A1:L{worksheet.max_row}"

    if "ProcessingRecords" in worksheet.tables:
        worksheet.tables["ProcessingRecords"].ref = table_ref
        return

    tracking_table = Table(
        displayName="ProcessingRecords",
        ref=table_ref,
    )
    tracking_table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    worksheet.add_table(tracking_table)


def append_document_record(
    workbook_path: Path,
    source_path: Path,
    baidu_markdown_path: Path | None,
    restored_markdown_path: Path | None,
    status: str,
    error_message: str = "",
) -> Path:
    """追加单份简历的文档解析阶段结果。"""

    workbook_path.parent.mkdir(parents=True, exist_ok=True)
    workbook, worksheet = _open_or_create_workbook(workbook_path)

    worksheet.append(
        (
            datetime.now(),
            str(source_path.resolve()),
            _existing_path(baidu_markdown_path),
            _existing_path(restored_markdown_path),
            status,
            error_message,
            "",
            "待提取" if status == "成功" else "",
            "",
            "",
            "",
            "",
        )
    )

    current_row = worksheet.max_row
    worksheet.cell(current_row, 1).number_format = (
        "yyyy-mm-dd hh:mm:ss"
    )

    for cell in worksheet[current_row]:
        cell.alignment = Alignment(
            vertical="top",
            wrap_text=True,
        )

    _set_status_style(
        worksheet.cell(current_row, 5),
        status,
    )
    _set_status_style(
        worksheet.cell(current_row, 8),
        worksheet.cell(current_row, 8).value or "",
    )
    _set_status_style(
        worksheet.cell(current_row, 11),
        worksheet.cell(current_row, 11).value or "",
    )
    _refresh_table(worksheet)
    workbook.save(workbook_path)
    workbook.close()
    return workbook_path.resolve()


def update_extraction_record(
    workbook_path: Path,
    source_path: Path,
    restored_markdown_path: Path,
    json_path: Path | None,
    status: str,
    error_message: str = "",
) -> Path:
    """更新最近一次匹配记录的信息提取阶段状态。"""

    workbook, worksheet = _open_or_create_workbook(workbook_path)
    source_value = str(source_path.resolve())
    restored_value = str(restored_markdown_path.resolve())
    target_row = None

    for row in range(worksheet.max_row, 1, -1):
        if (
            worksheet.cell(row, 2).value == source_value
            and worksheet.cell(row, 4).value == restored_value
        ):
            target_row = row
            break

    if target_row is None:
        workbook.close()
        raise RuntimeError(
            "处理记录中找不到对应的文档解析成功记录。"
        )

    worksheet.cell(
        target_row,
        7,
        _existing_path(json_path),
    )
    worksheet.cell(target_row, 8, status)
    worksheet.cell(target_row, 9, error_message)
    worksheet.cell(
        target_row,
        11,
        "待生成" if status == "成功" else "",
    )
    _set_status_style(
        worksheet.cell(target_row, 8),
        status,
    )
    _set_status_style(
        worksheet.cell(target_row, 11),
        worksheet.cell(target_row, 11).value or "",
    )
    _refresh_table(worksheet)
    workbook.save(workbook_path)
    workbook.close()
    return workbook_path.resolve()


def update_generation_record(
    workbook_path: Path,
    source_path: Path,
    json_path: Path,
    docx_path: Path | None,
    status: str,
    error_message: str = "",
) -> Path:
    """更新最近一次匹配记录的标准Word生成阶段状态。"""

    workbook, worksheet = _open_or_create_workbook(workbook_path)
    source_value = str(source_path.resolve())
    json_value = str(json_path.resolve())
    target_row = None

    for row in range(worksheet.max_row, 1, -1):
        if (
            worksheet.cell(row, 2).value == source_value
            and worksheet.cell(row, 7).value == json_value
        ):
            target_row = row
            break

    if target_row is None:
        workbook.close()
        raise RuntimeError(
            "处理记录中找不到对应的信息提取成功记录。"
        )

    worksheet.cell(target_row, 10, _existing_path(docx_path))
    worksheet.cell(target_row, 11, status)
    worksheet.cell(target_row, 12, error_message)
    _set_status_style(worksheet.cell(target_row, 11), status)
    _refresh_table(worksheet)
    workbook.save(workbook_path)
    workbook.close()
    return workbook_path.resolve()


def _existing_path(path: Path | None) -> str:
    """仅返回实际存在文件的绝对路径。"""

    if path and path.exists():
        return str(path.resolve())

    return ""
