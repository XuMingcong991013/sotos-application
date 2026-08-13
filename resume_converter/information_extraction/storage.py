"""
结构化简历JSON与汇总Excel的持久化工具。
"""

from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo


SUMMARY_WORKBOOK_NAME = "resume_information.xlsx"

SHEET_DEFINITIONS = {
    "候选人汇总": (
        "候选人ID",
        "姓名",
        "性别",
        "出生年份",
        "籍贯",
        "专业技能",
        "技能来源",
        "工作年限",
        "工作年限依据",
        "原始文件路径",
        "恢复Markdown路径",
        "JSON路径",
        "提取问题数",
    ),
    "工作经历": (
        "候选人ID",
        "姓名",
        "时间",
        "开始时间",
        "结束时间",
        "公司名称",
        "岗位名称",
        "工作描述",
        "JSON路径",
    ),
    "项目经历": (
        "候选人ID",
        "姓名",
        "项目名称",
        "岗位名称",
        "时间",
        "开始时间",
        "结束时间",
        "项目描述",
        "JSON路径",
    ),
    "教育经历": (
        "候选人ID",
        "姓名",
        "时间",
        "原始时间",
        "开始时间",
        "结束时间",
        "学校",
        "学历",
        "专业",
        "日期处理说明",
        "JSON路径",
    ),
    "提取问题": (
        "候选人ID",
        "姓名",
        "字段",
        "状态",
        "原文依据",
        "问题说明",
        "JSON路径",
    ),
}


def save_extracted_json(
    data: dict[str, Any],
    output_path: Path,
) -> Path:
    """以UTF-8保存单份候选人的完整结构化过程数据。"""

    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(
        json.dumps(
            data,
            ensure_ascii=False,
            indent=2,
        ) + "\n",
        encoding="utf-8",
    )
    return output_path.resolve()


def rebuild_summary_workbook(
    extraction_directory: Path,
) -> Path:
    """从目录中的全部JSON重建可筛选的多工作表汇总。"""

    json_files = sorted(
        extraction_directory.glob("*_extracted.json")
    )
    records = []

    for json_path in json_files:
        data = json.loads(
            json_path.read_text(encoding="utf-8")
        )
        data["_json_path"] = str(json_path.resolve())
        records.append(data)

    workbook = Workbook()
    workbook.remove(workbook.active)
    worksheets = {}

    for sheet_name, headers in SHEET_DEFINITIONS.items():
        worksheet = workbook.create_sheet(sheet_name)
        worksheet.append(headers)
        worksheets[sheet_name] = worksheet

    for data in records:
        _append_record(worksheets, data)

    for index, (sheet_name, worksheet) in enumerate(
        worksheets.items(),
        start=1,
    ):
        _style_sheet(worksheet, sheet_name)
        _add_table(worksheet, f"ResumeTable{index}")

    workbook_path = (
        extraction_directory / SUMMARY_WORKBOOK_NAME
    )
    workbook.save(workbook_path)
    workbook.close()
    return workbook_path.resolve()


def _append_record(worksheets, data: dict[str, Any]) -> None:
    """把一份JSON拆分追加到五张业务工作表。"""

    candidate_id = data.get("candidate_id", "")
    basic = data.get("basic_information", {})
    name = basic.get("name", "")
    json_path = data.get("_json_path", "")
    work_years = data.get("work_years", {})
    work_years_note = "；".join(
        value
        for value in (
            work_years.get("evidence", ""),
            work_years.get("note", ""),
        )
        if value
    )

    worksheets["候选人汇总"].append(
        (
            candidate_id,
            name,
            basic.get("gender", ""),
            basic.get("birth_year", ""),
            basic.get("native_place", ""),
            data.get("professional_skills", ""),
            data.get("professional_skills_source", ""),
            work_years.get("value"),
            work_years_note,
            data.get("source_file", ""),
            data.get("restored_markdown_file", ""),
            json_path,
            len(data.get("extraction_issues", [])),
        )
    )

    for item in data.get("work_experiences", []):
        worksheets["工作经历"].append(
            (
                candidate_id,
                name,
                item.get("time", ""),
                item.get("start_date", ""),
                item.get("end_date", ""),
                item.get("company_name", ""),
                item.get("position_name", ""),
                item.get("description", ""),
                json_path,
            )
        )

    for item in data.get("project_experiences", []):
        worksheets["项目经历"].append(
            (
                candidate_id,
                name,
                item.get("project_name", ""),
                item.get("position_name", ""),
                item.get("time", ""),
                item.get("start_date", ""),
                item.get("end_date", ""),
                item.get("description", ""),
                json_path,
            )
        )

    for item in data.get("education_experiences", []):
        worksheets["教育经历"].append(
            (
                candidate_id,
                name,
                item.get("time", ""),
                item.get("original_time", ""),
                item.get("start_date", ""),
                item.get("end_date", ""),
                item.get("school_name", ""),
                item.get("degree", ""),
                item.get("major", ""),
                item.get("normalization_note", ""),
                json_path,
            )
        )

    for issue in data.get("extraction_issues", []):
        worksheets["提取问题"].append(
            (
                candidate_id,
                name,
                issue.get("field", ""),
                issue.get("status", ""),
                issue.get("evidence", ""),
                issue.get("note", ""),
                json_path,
            )
        )


def _style_sheet(worksheet, sheet_name: str) -> None:
    """应用适合长文本和批量筛选的统一样式。"""

    worksheet.freeze_panes = "A2"
    worksheet.sheet_view.showGridLines = False
    worksheet.auto_filter.ref = worksheet.dimensions
    worksheet.row_dimensions[1].height = 24
    header_fill = PatternFill(
        fill_type="solid",
        fgColor="1F4E78",
    )

    for cell in worksheet[1]:
        cell.font = Font(
            color="FFFFFF",
            bold=True,
        )
        cell.fill = header_fill
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True,
            )

    default_widths = {
        "候选人汇总": [20, 12, 10, 12, 16, 48, 12, 12, 42, 42, 42, 42, 12],
        "工作经历": [20, 12, 20, 12, 12, 28, 24, 56, 42],
        "项目经历": [20, 12, 28, 24, 20, 12, 12, 56, 42],
        "教育经历": [20, 12, 20, 20, 12, 12, 28, 12, 24, 42, 42],
        "提取问题": [20, 12, 34, 12, 48, 52, 42],
    }

    for index, width in enumerate(
        default_widths[sheet_name],
        start=1,
    ):
        worksheet.column_dimensions[
            worksheet.cell(1, index).column_letter
        ].width = width

    if sheet_name == "候选人汇总":
        for row in range(2, worksheet.max_row + 1):
            worksheet.cell(row, 8).number_format = "0.0"
            worksheet.cell(row, 13).number_format = "0"


def _add_table(worksheet, table_name: str) -> None:
    """有数据行时添加Excel结构化表。"""

    if worksheet.max_row < 2:
        return

    table = Table(
        displayName=table_name,
        ref=worksheet.dimensions,
    )
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    worksheet.add_table(table)
