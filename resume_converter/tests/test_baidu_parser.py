"""百度办公文档识别和 DOCX 本地解析的离线测试。"""

from __future__ import annotations

import io
import tempfile
import unittest
from contextlib import redirect_stdout
from pathlib import Path
from unittest.mock import Mock, patch

import pymupdf
from docx import Document

from utils import baidu_parser


class OfficeOcrMarkdownTests(unittest.TestCase):
    """验证办公文档识别 JSON 到 Markdown 的稳定转换。"""

    def test_layout_headings_and_page_attributes_are_handled(self) -> None:
        """标题保留语义，页眉页脚和装饰图片被忽略。"""

        result = {
            "results": [
                {"words": {"word": "重复页眉"}},
                {"words": {"word": "张三"}},
                {"words": {"word": "工作经历"}},
                {"words": {"word": "2020-至今 示例公司"}},
                {"words": {"word": "LOGO"}},
                {"words": {"word": "第1页"}},
            ],
            "sections": [
                {"attribute": "header", "sec_idx": {"idx": "0"}},
                {"attribute": "number", "sec_idx": [{"idx": 5}]},
            ],
            "layouts": [
                {"layout": "doc_title", "layout_idx": [1]},
                {"layout": "title", "layout_idx": {"idx": "2"}},
                {"layout": "figure", "layout_idx": [4]},
            ],
        }

        markdown = baidu_parser.office_ocr_result_to_markdown(result)

        self.assertEqual(
            markdown,
            "# 张三\n\n## 工作经历\n\n2020-至今 示例公司\n",
        )

    def test_response_without_body_is_rejected(self) -> None:
        """成功响应没有可用正文时给出明确错误。"""

        with self.assertRaisesRegex(RuntimeError, "没有返回可用正文"):
            baidu_parser.office_ocr_result_to_markdown(
                {"results": [{"words": {"word": ""}}]}
            )

    def test_empty_page_can_be_skipped_by_document_parser(self) -> None:
        """多页文档中的空白页允许返回空字符串供上层跳过。"""

        markdown = baidu_parser.office_ocr_result_to_markdown(
            {"results": [{"words": {"word": ""}}]},
            allow_empty=True,
        )

        self.assertEqual(markdown, "")


class OfficeOcrRequestTests(unittest.TestCase):
    """验证后缀分流和请求参数，不访问百度服务。"""

    def test_image_uses_office_document_recognition(self) -> None:
        """图片只调用办公文档识别，并写出中间 Markdown。"""

        with tempfile.TemporaryDirectory() as temp:
            temp_path = Path(temp)
            image_path = temp_path / "resume.png"
            output_dir = temp_path / "output"
            image_path.write_bytes(b"fake-image")

            token_response = Mock()
            token_response.raise_for_status.return_value = None
            token_response.json.return_value = {"access_token": "test-token"}

            ocr_response = Mock()
            ocr_response.raise_for_status.return_value = None
            ocr_response.json.return_value = {
                "log_id": "123",
                "results": [{"words": {"word": "候选人姓名"}}],
            }

            with (
                patch.dict(
                    baidu_parser.os.environ,
                    {
                        "BAIDU_API_KEY": "test-key",
                        "BAIDU_SECRET_KEY": "test-secret",
                    },
                ),
                patch.object(
                    baidu_parser.requests,
                    "post",
                    side_effect=[token_response, ocr_response],
                ) as post,
                redirect_stdout(io.StringIO()),
            ):
                result = baidu_parser.BaiduParser(
                    str(image_path),
                    str(output_dir),
                )

            self.assertEqual(post.call_count, 2)
            self.assertEqual(post.call_args_list[1].args[0], baidu_parser.OFFICE_OCR_URL)
            request_data = post.call_args_list[1].kwargs["data"]
            self.assertIn("image", request_data)
            self.assertNotIn("pdf_file", request_data)
            self.assertEqual(request_data["layout_analysis"], "true")
            self.assertEqual(
                Path(result).read_text(encoding="utf-8"),
                "候选人姓名\n",
            )

    def test_pdf_is_recognized_one_page_per_request(self) -> None:
        """多页 PDF 按页调用，并传递从 1 开始的页码。"""

        with tempfile.TemporaryDirectory() as temp:
            temp_path = Path(temp)
            pdf_path = temp_path / "resume.pdf"
            document = pymupdf.open()
            document.new_page()
            document.new_page()
            document.save(pdf_path)
            document.close()

            responses = [
                {"results": [{"words": {"word": "第一页"}}]},
                {"results": [{"words": {"word": "第二页"}}]},
            ]

            with (
                patch.object(
                    baidu_parser,
                    "get_access_token",
                    return_value="test-token",
                ),
                patch.object(
                    baidu_parser,
                    "_request_office_ocr",
                    side_effect=responses,
                ) as request_ocr,
                patch.dict(
                    baidu_parser.os.environ,
                    {
                        "BAIDU_API_KEY": "test-key",
                        "BAIDU_SECRET_KEY": "test-secret",
                    },
                ),
                redirect_stdout(io.StringIO()),
            ):
                result = baidu_parser.BaiduParser(
                    str(pdf_path),
                    str(temp_path / "output"),
                )

            self.assertEqual(request_ocr.call_count, 2)
            self.assertEqual(
                [call.kwargs["page_number"] for call in request_ocr.call_args_list],
                [1, 2],
            )
            self.assertEqual(
                Path(result).read_text(encoding="utf-8"),
                "第一页\n\n第二页\n",
            )

    def test_ocr_error_is_not_retried_or_sent_to_another_service(self) -> None:
        """配额错误直接上报，不自动切换到其他付费接口。"""

        response = Mock()
        response.raise_for_status.return_value = None
        response.json.return_value = {
            "error_code": 17,
            "error_msg": "Open api daily request limit reached",
        }

        with (
            patch.object(baidu_parser.requests, "post", return_value=response) as post,
            self.assertRaisesRegex(RuntimeError, "办公文档识别失败"),
        ):
            baidu_parser._request_office_ocr(
                access_token="test-token",
                file_field="image",
                file_base64="YWJj",
            )

        self.assertEqual(post.call_count, 1)


class LocalDocxParserTests(unittest.TestCase):
    """验证 DOCX 完全在本地解析。"""

    def test_docx_preserves_headings_lists_and_tables(self) -> None:
        """段落、标题、列表和表格按原顺序进入中间 Markdown。"""

        with tempfile.TemporaryDirectory() as temp:
            temp_path = Path(temp)
            docx_path = temp_path / "resume.docx"
            document = Document()
            document.add_heading("个人简历", level=1)
            document.add_paragraph("熟悉 Python", style="List Bullet")
            table = document.add_table(rows=1, cols=2)
            table.cell(0, 0).text = "公司"
            table.cell(0, 1).text = "岗位"
            document.save(docx_path)

            with (
                patch.object(baidu_parser.requests, "post") as post,
                redirect_stdout(io.StringIO()),
            ):
                result = baidu_parser.BaiduParser(
                    str(docx_path),
                    str(temp_path / "output"),
                )

            post.assert_not_called()
            self.assertEqual(
                Path(result).read_text(encoding="utf-8"),
                "# 个人简历\n\n- 熟悉 Python\n\n公司 | 岗位\n",
            )

    def test_empty_docx_still_generates_clear_markdown(self) -> None:
        """空 DOCX 生成明确提示，不因缺少参考文本提前失败。"""

        with tempfile.TemporaryDirectory() as temp:
            temp_path = Path(temp)
            docx_path = temp_path / "empty.docx"
            Document().save(docx_path)

            result = baidu_parser.BaiduParser(
                str(docx_path),
                str(temp_path / "output"),
            )

            self.assertEqual(
                Path(result).read_text(encoding="utf-8"),
                baidu_parser.EMPTY_DOCX_MARKDOWN,
            )


if __name__ == "__main__":
    unittest.main()
