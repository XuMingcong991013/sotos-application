"""
简历转换流程的离线单元测试。

所有百度和LLM调用均使用模拟对象，测试不会访问外部服务。
"""

from __future__ import annotations

import io
import tempfile
import unittest
from contextlib import redirect_stdout
from pathlib import Path
from unittest.mock import patch

import pymupdf
from docx import Document
from openpyxl import load_workbook

import data_parser
from utils import restore_markdown


class DataParserTests(unittest.TestCase):
    """验证统一入口的路径、后缀和错误处理。"""

    def test_supported_suffixes_use_stem_for_all_output_paths(self) -> None:
        """五类文件均使用原文件stem拼接中间和最终路径。"""

        for suffix in sorted(restore_markdown.SUPPORTED_SUFFIXES):
            with self.subTest(suffix=suffix), tempfile.TemporaryDirectory(
                dir=Path.cwd()
            ) as temp:
                temp_path = Path(temp)
                source_path = temp_path / f"candidate.resume{suffix.upper()}"
                output_dir = temp_path / "outputs"
                process_dir = (
                    output_dir
                    / "process_data"
                    / "document_parsing"
                )
                source_path.write_bytes(b"test")

                def fake_baidu_parser(
                    input_file_path: str,
                    output_dir: str,
                ) -> str:
                    generated_path = (
                        Path(output_dir)
                        / f"{Path(input_file_path).stem}.md"
                    )
                    generated_path.write_text(
                        "百度内容",
                        encoding="utf-8",
                    )
                    return str(generated_path.resolve())

                def fake_restore_markdown(
                    input_file_path: str,
                    input_file: str,
                    output_file: str,
                ) -> str:
                    self.assertEqual(
                        Path(input_file),
                        process_dir / "candidate.resume.md",
                    )
                    self.assertEqual(
                        Path(output_file),
                        process_dir / "candidate.resume_restored.md",
                    )
                    return str(Path(output_file).resolve())

                with (
                    patch.object(
                        data_parser,
                        "BaiduParser",
                        side_effect=fake_baidu_parser,
                    ),
                    patch.object(
                        data_parser,
                        "RestoreMarkdown",
                        side_effect=fake_restore_markdown,
                    ),
                    redirect_stdout(io.StringIO()),
                ):
                    result = data_parser.DataParser(
                        str(source_path),
                        str(output_dir),
                    )

                self.assertEqual(
                    result,
                    str(
                        (
                            output_dir
                            / "process_data"
                            / "document_parsing"
                            / "candidate.resume_restored.md"
                        ).resolve()
                    ),
                )

    def test_error_log_is_appended_and_returns_none(self) -> None:
        """同一天的真正异常追加到同一个日志并返回None。"""

        with tempfile.TemporaryDirectory(
            dir=Path.cwd()
        ) as temp:
            temp_path = Path(temp)
            output_dir = temp_path / "outputs"
            process_dir = output_dir / "process_data"
            missing_path = temp_path / "missing.pdf"

            with redirect_stdout(io.StringIO()):
                first_result = data_parser.DataParser(
                    str(missing_path),
                    str(output_dir),
                )
                second_result = data_parser.DataParser(
                    str(missing_path),
                    str(output_dir),
                )

            error_files = list(
                process_dir.glob("error_*.txt")
            )

            self.assertIsNone(first_result)
            self.assertIsNone(second_result)
            self.assertEqual(len(error_files), 1)

            log_content = error_files[0].read_text(
                encoding="utf-8"
            )
            self.assertEqual(
                log_content.splitlines().count("=" * 80),
                2,
            )
            self.assertEqual(
                sum(
                    line.startswith("输入文件：")
                    for line in log_content.splitlines()
                ),
                2,
            )
            self.assertIn("Traceback", log_content)

            workbook = load_workbook(
                process_dir / "processing_records.xlsx",
                data_only=True,
            )
            worksheet = workbook["处理记录"]
            self.assertEqual(worksheet.max_row, 3)
            self.assertEqual(worksheet["E2"].value, "失败")
            self.assertEqual(worksheet["E3"].value, "失败")
            workbook.close()

    def test_doc_is_rejected_before_baidu_api_call(self) -> None:
        """旧版DOC会被明确拒绝，不会调用百度API。"""

        with tempfile.TemporaryDirectory(
            dir=Path.cwd()
        ) as temp:
            temp_path = Path(temp)
            source_path = temp_path / "resume.doc"
            output_dir = temp_path / "outputs"
            source_path.write_bytes(b"legacy-doc")

            with (
                patch.object(
                    data_parser,
                    "BaiduParser",
                ) as baidu_parser,
                redirect_stdout(io.StringIO()) as stdout,
            ):
                result = data_parser.DataParser(
                    str(source_path),
                    str(output_dir),
                )

            self.assertIsNone(result)
            baidu_parser.assert_not_called()
            self.assertIn(
                "不支持旧版DOC文件，请先转换为DOCX格式",
                stdout.getvalue(),
            )
            self.assertEqual(
                len(
                    list(
                        (
                            output_dir
                            / "process_data"
                        ).glob("error_*.txt")
                    )
                ),
                1,
            )

    def test_success_record_contains_process_file_paths(self) -> None:
        """成功记录包含原文件和两个实际过程文件的绝对路径。"""

        with tempfile.TemporaryDirectory(
            dir=Path.cwd()
        ) as temp:
            temp_path = Path(temp)
            source_path = temp_path / "resume.pdf"
            output_dir = temp_path / "outputs"
            process_root = output_dir / "process_data"
            process_dir = process_root / "document_parsing"
            source_path.write_bytes(b"pdf")

            def fake_baidu_parser(
                input_file_path: str,
                output_dir: str,
            ) -> str:
                markdown_path = Path(output_dir) / "resume.md"
                markdown_path.write_text("百度", encoding="utf-8")
                return str(markdown_path.resolve())

            def fake_restore_markdown(
                input_file_path: str,
                input_file: str,
                output_file: str,
            ) -> str:
                restored_path = Path(output_file)
                restored_path.write_text("恢复", encoding="utf-8")
                return str(restored_path.resolve())

            with (
                patch.object(
                    data_parser,
                    "BaiduParser",
                    side_effect=fake_baidu_parser,
                ),
                patch.object(
                    data_parser,
                    "RestoreMarkdown",
                    side_effect=fake_restore_markdown,
                ),
                redirect_stdout(io.StringIO()),
            ):
                result = data_parser.DataParser(
                    str(source_path),
                    str(output_dir),
                )

            workbook_path = (
                process_root / "processing_records.xlsx"
            )
            workbook = load_workbook(
                workbook_path,
                data_only=True,
            )
            worksheet = workbook["处理记录"]

            self.assertEqual(
                result,
                str(
                    (
                        process_dir
                        / "resume_restored.md"
                    ).resolve()
                ),
            )
            self.assertEqual(worksheet.max_row, 2)
            self.assertEqual(
                worksheet["B2"].value,
                str(source_path.resolve()),
            )
            self.assertEqual(
                worksheet["C2"].value,
                str((process_dir / "resume.md").resolve()),
            )
            self.assertEqual(
                worksheet["D2"].value,
                str(
                    (
                        process_dir
                        / "resume_restored.md"
                    ).resolve()
                ),
            )
            self.assertEqual(worksheet["E2"].value, "成功")
            self.assertEqual(worksheet.freeze_panes, "A2")
            self.assertIn("ProcessingRecords", worksheet.tables)
            workbook.close()


class ReferenceTextTests(unittest.TestCase):
    """验证不同文件后缀的辅助参考文本策略。"""

    def test_scanned_pdf_without_text_layer_uses_clear_prompt(self) -> None:
        """空文本层PDF不报错，并返回明确的无参考提示。"""

        with tempfile.TemporaryDirectory(
            dir=Path.cwd()
        ) as temp:
            pdf_path = Path(temp) / "scan.pdf"
            document = pymupdf.open()
            document.new_page()
            document.save(pdf_path)
            document.close()

            reference_text = (
                restore_markdown.extract_reference_text(
                    pdf_path
                )
            )

            self.assertEqual(
                reference_text,
                restore_markdown.NO_REFERENCE_TEXT[".pdf"],
            )

    def test_docx_extracts_paragraphs_and_tables(self) -> None:
        """DOCX参考文本包含段落和表格文字。"""

        with tempfile.TemporaryDirectory(
            dir=Path.cwd()
        ) as temp:
            docx_path = Path(temp) / "resume.docx"
            document = Document()
            document.add_paragraph("个人简介")
            table = document.add_table(rows=1, cols=2)
            table.cell(0, 0).text = "公司"
            table.cell(0, 1).text = "岗位"
            document.add_paragraph("项目经历")
            document.save(docx_path)

            reference_text = (
                restore_markdown.extract_reference_text(
                    docx_path
                )
            )

            self.assertEqual(
                reference_text.splitlines(),
                [
                    "个人简介",
                    "公司 | 岗位",
                    "项目经历",
                ],
            )

    def test_images_use_no_reference_prompt(self) -> None:
        """图片不引入本地OCR。"""

        for suffix in (".jpg", ".jpeg", ".png"):
            with self.subTest(suffix=suffix):
                source_path = Path(f"resume{suffix}")
                self.assertEqual(
                    restore_markdown.extract_reference_text(
                        source_path
                    ),
                    restore_markdown.NO_REFERENCE_TEXT[suffix],
                )


class RestoreMarkdownTests(unittest.TestCase):
    """验证结构恢复警告不会中断最终文件生成。"""

    def test_validation_warning_does_not_block_output(self) -> None:
        """即使自动校验告警，仍写文件并返回绝对路径。"""

        with tempfile.TemporaryDirectory(
            dir=Path.cwd()
        ) as temp:
            temp_path = Path(temp)
            source_path = temp_path / "resume.jpg"
            input_path = temp_path / "resume.md"
            output_path = temp_path / "resume_restored.md"
            source_path.write_bytes(b"not-a-real-image")
            input_path.write_text(
                "手机号：13800138000\n" + "正文" * 20,
                encoding="utf-8",
            )

            with (
                patch.dict(
                    restore_markdown.os.environ,
                    {
                        "LLM_API_KEY": "offline-test-key",
                        "LLM_BASE_URL": "http://127.0.0.1:1/v1",
                        "LLM_MODEL": "offline-test-model",
                    },
                    clear=False,
                ),
                patch.object(
                    restore_markdown,
                    "load_dotenv",
                ),
                patch.object(
                    restore_markdown,
                    "call_llm_restore",
                    return_value="# 恢复结果\n",
                ),
                redirect_stdout(io.StringIO()) as stdout,
            ):
                result = restore_markdown.RestoreMarkdown(
                    str(source_path),
                    str(input_path),
                    str(output_path),
                )

            self.assertEqual(
                result,
                str(output_path.resolve()),
            )
            self.assertEqual(
                output_path.read_text(encoding="utf-8"),
                "# 恢复结果\n",
            )
            self.assertIn("警告", stdout.getvalue())
            self.assertIn("不会中断文件生成", stdout.getvalue())
