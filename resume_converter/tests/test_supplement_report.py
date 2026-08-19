"""
最终简历待补充信息清单的离线测试。
"""

from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from resume_generation.supplement_report import (
    SUPPLEMENT_HEADERS,
    SUPPLEMENT_SHEET_NAME,
    collect_chery_supplement_items,
    collect_supplement_items,
    update_supplement_report,
)


class SupplementReportTests(unittest.TestCase):
    """验证必填规则、选填排除和同路径更新。"""

    def test_chery_collects_extended_fields_and_ai_confirmation(self) -> None:
        """奇瑞清单包含联系方式、项目业绩和AI内容确认。"""

        data = {
            "basic_information": {
                "name": "李四",
                "gender": "女",
                "birth_date": "1995年6月",
                "phone": "",
                "email": "",
                "native_place": "江苏南京",
            },
            "work_years": {"value": 6.0},
            "self_evaluation": "AI评价",
            "self_evaluation_source": "generated",
            "professional_skills": "原文技能",
            "professional_skills_source": "explicit",
            "work_experiences": [
                {
                    "time": "2020/07-至今",
                    "company_name": "甲公司",
                    "position_name": "工程师",
                }
            ],
            "project_experiences": [
                {
                    "project_name": "项目A",
                    "description": "项目描述",
                    "achievement": "",
                }
            ],
            "education_experiences": [
                {
                    "time": "2016/09-2020/06",
                    "school_name": "某大学",
                    "major": "电子信息工程",
                    "degree": "本科",
                }
            ],
        }

        self.assertEqual(
            collect_chery_supplement_items(data),
            [
                "请补充联系电话",
                "请补充邮箱",
                "请人工确认AI生成的自我评价",
                "请补充项目1的工作业绩",
            ],
        )

    def test_collects_only_required_missing_information(self) -> None:
        """工作描述及项目岗位、时间缺失时不应要求补充。"""

        data = {
            "basic_information": {
                "name": "张三",
                "gender": "男",
                "birth_year": "",
                "native_place": "湖北武汉",
            },
            "professional_skills": "熟悉 Python",
            "work_experiences": [
                {
                    "time": "2024/01-至今",
                    "company_name": "甲公司",
                    "position_name": "工程师",
                    "description": "",
                }
            ],
            "project_experiences": [
                {
                    "project_name": "项目A",
                    "position_name": "",
                    "time": "",
                    "description": "",
                }
            ],
            "education_experiences": [
                {
                    "time": "2020/09-2024/06",
                    "school_name": "某大学",
                    "degree": "本科",
                    "major": "计算机科学与技术",
                }
            ],
            "work_years": {"value": 2.0},
        }

        items = collect_supplement_items(data)

        self.assertEqual(
            items,
            [
                "请补充出生年份",
                "请确认岗位职称",
                "请补充项目一的项目描述",
            ],
        )
        joined = "\n".join(items)
        self.assertNotIn("工作描述", joined)
        self.assertNotIn("项目岗位", joined)
        self.assertNotIn("项目时间", joined)

    def test_merges_missing_education_and_work_years_request(self) -> None:
        """整段教育缺失时把工作年限原因合并到同一事项。"""

        data = {
            "basic_information": {
                "name": "李四",
                "gender": "女",
                "birth_year": "1998",
                "native_place": "安徽合肥",
            },
            "professional_skills": "硬件设计",
            "work_experiences": [],
            "project_experiences": [],
            "education_experiences": [],
            "work_years": {"value": None},
        }

        items = collect_supplement_items(data)
        education_items = [item for item in items if "教育经历" in item]

        self.assertEqual(len(education_items), 1)
        self.assertEqual(
            education_items[0],
            "请补充教育经历：就读时间、学校、学历、专业；"
            "补充毕业时间后才能计算工作年限",
        )
        self.assertEqual(
            sum("工作年限" in item for item in items),
            1,
        )

    def test_creates_and_updates_single_row_by_final_path(self) -> None:
        """重复生成同一路径时更新原记录而不追加重复行。"""

        with tempfile.TemporaryDirectory(dir=Path.cwd()) as temp:
            directory = Path(temp)
            final_path = directory / "张三_标准简历.docx"
            final_path.write_bytes(b"docx")
            workbook_path = directory / "待补充信息.xlsx"
            data = {
                "basic_information": {
                    "name": "张三",
                    "gender": "男",
                    "birth_year": "1999",
                    "native_place": "湖北武汉",
                },
                "professional_skills": "Python",
                "work_experiences": [
                    {
                        "time": "2024/01-至今",
                        "company_name": "甲公司",
                        "position_name": "工程师",
                    }
                ],
                "project_experiences": [
                    {
                        "project_name": "项目A",
                        "description": "负责接口开发",
                    }
                ],
                "education_experiences": [
                    {
                        "time": "2020/09-2024/06",
                        "school_name": "某大学",
                        "degree": "本科",
                        "major": "计算机",
                    }
                ],
                "work_years": {"value": 2.0},
            }

            update_supplement_report(data, final_path, workbook_path)
            data["basic_information"]["name"] = "张三（已确认）"
            update_supplement_report(data, final_path, workbook_path)

            workbook = load_workbook(workbook_path, data_only=True)
            worksheet = workbook[SUPPLEMENT_SHEET_NAME]
            self.assertEqual(worksheet.max_row, 2)
            self.assertEqual(
                tuple(cell.value for cell in worksheet[1]),
                SUPPLEMENT_HEADERS,
            )
            self.assertEqual(worksheet["A2"].value, "张三（已确认）")
            self.assertEqual(
                worksheet["B2"].value,
                str(final_path.resolve()),
            )
            self.assertEqual(
                worksheet["C2"].value,
                "1. 请确认岗位职称",
            )
            self.assertEqual(
                worksheet.tables["SupplementInformation"].ref,
                "A1:C2",
            )
            workbook.close()
