"""
标准简历Word生成模块的离线测试。
"""

from __future__ import annotations

import io
import json
import tempfile
import unittest
from contextlib import redirect_stdout
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.oxml.ns import qn
from lxml import etree
from openpyxl import load_workbook

from resume_generation import ResumeGenerator
from resume_generation.generator import (
    FONT_NAME,
    PLACEHOLDER_COLOR,
    PROJECT_CONTENT_SPACE_BEFORE,
    _xml_text,
)
from utils.processing_records import (
    append_document_record,
    update_extraction_record,
)


def sample_data(source_path: Path) -> dict:
    """返回覆盖全部Word分区的结构化测试数据。"""

    return {
        "source_file": str(source_path.resolve()),
        "basic_information": {
            "name": "张三",
            "gender": "男",
            "birth_year": "1999",
            "native_place": "湖北武汉",
        },
        "professional_skills": (
            "1. 熟悉 Python、FastAPI\n"
            "· 掌握 MySQL、Redis"
        ),
        "work_experiences": [
            {
                "time": "2024/01-至今",
                "company_name": "甲公司",
                "position_name": "高级后端工程师",
                "description": "负责内部接口开发。",
            },
            {
                "time": "2022/07-2023/12",
                "company_name": "乙公司",
                "position_name": "开发工程师",
                "description": "负责业务系统维护。",
            },
        ],
        "project_experiences": [
            {
                "project_name": "项目A",
                "position_name": "项目负责人",
                "time": "2025/01-2025/06",
                "description": "负责项目A接口设计和交付。",
            },
            {
                "project_name": "项目B",
                "position_name": "",
                "time": "",
                "description": "",
            },
        ],
        "education_experiences": [
            {
                "time": "2022/09-2026/06",
                "school_name": "某大学",
                "degree": "本科",
                "major": "计算机科学与技术",
            }
        ],
        "work_years": {"value": 0.5},
    }


class ResumeGenerationTests(unittest.TestCase):
    """验证模板内容、两种基本资料模式及过程记录。"""

    def test_generates_no_photo_resume_and_updates_tracking(self) -> None:
        """无证件照模式生成两栏资料区并更新Word生成状态。"""

        with tempfile.TemporaryDirectory(dir=Path.cwd()) as temp:
            temp_path = Path(temp)
            output_dir = temp_path / "outputs"
            process_dir = output_dir / "process_data"
            parsing_dir = process_dir / "document_parsing"
            extraction_dir = process_dir / "information_extraction"
            parsing_dir.mkdir(parents=True)
            extraction_dir.mkdir(parents=True)
            source_path = temp_path / "resume.pdf"
            restored_path = parsing_dir / "resume_restored.md"
            json_path = extraction_dir / "resume_extracted.json"
            source_path.write_bytes(b"pdf")
            restored_path.write_text("简历", encoding="utf-8")
            json_path.write_text(
                json.dumps(sample_data(source_path), ensure_ascii=False),
                encoding="utf-8",
            )
            workbook_path = process_dir / "processing_records.xlsx"
            append_document_record(
                workbook_path,
                source_path,
                None,
                restored_path,
                "成功",
            )
            update_extraction_record(
                workbook_path,
                source_path,
                restored_path,
                json_path,
                "成功",
            )

            with redirect_stdout(io.StringIO()):
                result = ResumeGenerator(
                    str(json_path),
                    str(output_dir),
                    with_photo=False,
                )

            self.assertIsNotNone(result)
            result_path = Path(result)
            self.assertEqual(
                result_path.parent,
                (output_dir / "final_resumes").resolve(),
            )
            document = Document(result_path)
            self.assertEqual(len(document.tables[0].columns), 2)
            text = _all_document_text(document)
            self.assertIn("岗位职称：待补充", text)
            self.assertIn(
                "项目一：项目A | 项目负责人 | 2025/01-2025/06",
                text,
            )
            self.assertIn("项目二：项目B", text)
            self.assertIn("项目描述：待补充", text)
            self.assertNotIn("项目二：项目B | 岗位待补充", text)
            self.assertNotIn("项目二：项目B | 时间待补充", text)
            self.assertNotIn("项目一：项目A | 甲公司", text)
            self.assertIn("Project Experience", text)
            self.assertIn("2022/09-2026/06 | 某大学 | 本科 | 计算机科学与技术", text)
            self.assertIn("0.5年", text)
            project_paragraph = next(
                paragraph
                for paragraph in document.paragraphs
                if paragraph.text.startswith("项目一：")
            )
            self.assertEqual(
                project_paragraph.paragraph_format.space_before,
                PROJECT_CONTENT_SPACE_BEFORE,
            )
            _assert_all_declared_fonts_are_yahei(self, result_path)

            supplement_workbook = load_workbook(
                output_dir / "final_resumes" / "待补充信息.xlsx",
                data_only=True,
            )
            supplement_sheet = supplement_workbook["待补充信息"]
            self.assertEqual(supplement_sheet["A2"].value, "张三")
            self.assertEqual(
                supplement_sheet["B2"].value,
                str(result_path.resolve()),
            )
            self.assertIn(
                "请确认岗位职称",
                supplement_sheet["C2"].value,
            )
            self.assertIn(
                "请补充项目二的项目描述",
                supplement_sheet["C2"].value,
            )
            supplement_workbook.close()

            workbook = load_workbook(workbook_path, data_only=True)
            worksheet = workbook["处理记录"]
            self.assertEqual(worksheet["J2"].value, str(result_path.resolve()))
            self.assertEqual(worksheet["K2"].value, "成功")
            self.assertEqual(worksheet.max_column, 12)
            workbook.close()

    def test_photo_mode_uses_two_two_one_layout_and_placeholder(self) -> None:
        """证件照模式使用2:2:1三栏并且只显示占位。"""

        with tempfile.TemporaryDirectory(dir=Path.cwd()) as temp:
            temp_path = Path(temp)
            output_dir = temp_path / "outputs"
            source_path = temp_path / "resume.docx"
            json_path = temp_path / "resume_extracted.json"
            source_path.write_bytes(b"docx")
            json_path.write_text(
                json.dumps(sample_data(source_path), ensure_ascii=False),
                encoding="utf-8",
            )

            with redirect_stdout(io.StringIO()):
                result = ResumeGenerator(
                    str(json_path),
                    str(output_dir),
                    with_photo=True,
                )

            document = Document(result)
            basic_table = document.tables[0]
            self.assertEqual(len(basic_table.columns), 3)
            grid_widths = [
                int(column.get(qn("w:w")))
                for column in basic_table._tbl.tblGrid.gridCol_lst
            ]
            self.assertAlmostEqual(grid_widths[0] / grid_widths[2], 2, places=1)
            self.assertAlmostEqual(grid_widths[1] / grid_widths[2], 2, places=1)
            self.assertIn("证件照\n占位", basic_table.cell(0, 2).text)

    def test_missing_information_uses_stable_placeholders(self) -> None:
        """字段和整个分区缺失时保留可人工补充的稳定版式。"""

        with tempfile.TemporaryDirectory(dir=Path.cwd()) as temp:
            temp_path = Path(temp)
            output_dir = temp_path / "outputs"
            source_path = temp_path / "resume.pdf"
            json_path = temp_path / "resume_extracted.json"
            source_path.write_bytes(b"pdf")
            json_path.write_text(
                json.dumps(
                    {
                        "source_file": str(source_path.resolve()),
                        "basic_information": {
                            "name": "张三",
                            "gender": "",
                            "birth_year": "",
                            "native_place": "",
                        },
                        "professional_skills": "",
                        "work_experiences": [],
                        "project_experiences": [],
                        "education_experiences": [],
                        "work_years": {"value": None},
                    },
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )

            with redirect_stdout(io.StringIO()):
                result = ResumeGenerator(
                    str(json_path),
                    str(output_dir),
                )

            document = Document(result)
            text = _all_document_text(document)
            self.assertIn("性别：待补充", text)
            self.assertIn("出生年份：待补充", text)
            self.assertIn("籍贯：待补充", text)
            self.assertIn("岗位职称：待补充", text)
            self.assertIn("时间待补充", text)
            self.assertIn("公司名称待补充", text)
            self.assertIn("岗位待补充", text)
            self.assertNotIn("工作描述：待补充", text)
            self.assertIn("项目一：项目名称待补充", text)
            self.assertIn("项目描述：待补充", text)
            self.assertNotIn(
                "项目一：项目名称待补充 | 岗位待补充",
                text,
            )
            self.assertIn(
                "时间待补充 | 学校待补充 | 学历待补充 | 专业待补充",
                text,
            )

            placeholder_runs = [
                run
                for run in document._element.body.iter(qn("w:r"))
                if _xml_text(run) == "待补充"
            ]
            self.assertGreaterEqual(len(placeholder_runs), 10)
            self.assertTrue(
                all(
                    run.find(qn("w:rPr")).find(qn("w:color")).get(
                        qn("w:val")
                    )
                    == str(PLACEHOLDER_COLOR)
                    for run in placeholder_runs
                )
            )

            time_label_run = next(
                run
                for run in document._element.body.iter(qn("w:r"))
                if _xml_text(run) == "时间"
            )
            self.assertNotEqual(
                time_label_run.find(qn("w:rPr")).find(qn("w:color")).get(
                    qn("w:val")
                ),
                str(PLACEHOLDER_COLOR),
            )

    def test_missing_json_is_logged_and_returns_none(self) -> None:
        """真正异常写入统一错误日志并返回None。"""

        with tempfile.TemporaryDirectory(dir=Path.cwd()) as temp:
            output_dir = Path(temp) / "outputs"

            with redirect_stdout(io.StringIO()):
                result = ResumeGenerator(
                    str(Path(temp) / "missing.json"),
                    str(output_dir),
                )

            self.assertIsNone(result)
            error_file = next(
                (output_dir / "process_data").glob("error_*.txt")
            )
            error_text = error_file.read_text(encoding="utf-8")
            self.assertIn("处理阶段：标准简历生成", error_text)
            self.assertIn("找不到结构化简历JSON", error_text)

    def test_markdown_blank_lines_do_not_create_empty_word_paragraphs(self) -> None:
        """Markdown正文间的空行不应在Word中变成额外空段落。"""

        with tempfile.TemporaryDirectory(dir=Path.cwd()) as temp:
            temp_path = Path(temp)
            output_dir = temp_path / "outputs"
            source_path = temp_path / "resume.pdf"
            json_path = temp_path / "resume_extracted.json"
            source_path.write_bytes(b"pdf")
            data = sample_data(source_path)
            data["professional_skills"] = "技能第一行\n\n技能第二行"
            data["work_experiences"][0]["description"] = (
                "工作描述第一行\n\n工作描述第二行\n   \n工作描述第三行"
            )
            data["project_experiences"][0]["description"] = (
                "项目描述第一行\n\n项目描述第二行"
            )
            json_path.write_text(
                json.dumps(data, ensure_ascii=False),
                encoding="utf-8",
            )

            with redirect_stdout(io.StringIO()):
                result = ResumeGenerator(str(json_path), str(output_dir))

            document = Document(result)
            text_paragraphs = {
                paragraph.text: paragraph
                for paragraph in document.paragraphs
                if paragraph.text
            }
            self.assertIn("技能第一行", text_paragraphs)
            self.assertIn("技能第二行", text_paragraphs)
            self.assertIn("项目描述第一行", text_paragraphs)
            self.assertIn("项目描述第二行", text_paragraphs)

            work_table = document.tables[1]
            description_cell = work_table.cell(1, 0)
            self.assertEqual(
                [paragraph.text for paragraph in description_cell.paragraphs],
                [
                    "工作描述第一行",
                    "工作描述第二行",
                    "工作描述第三行",
                ],
            )

    def test_common_markdown_is_rendered_as_word_formatting(self) -> None:
        """三类长文本共用Markdown渲染，最终可见文本不泄漏控制标记。"""

        with tempfile.TemporaryDirectory(dir=Path.cwd()) as temp:
            temp_path = Path(temp)
            output_dir = temp_path / "outputs"
            source_path = temp_path / "resume.pdf"
            json_path = temp_path / "resume_extracted.json"
            source_path.write_bytes(b"pdf")
            data = sample_data(source_path)
            data["professional_skills"] = (
                "- 列表第一项\n"
                "- 列表第二项包含 C++、C#、char* 和 A*算法"
            )
            data["work_experiences"][0]["description"] = (
                "**通用标签：** 普通内容"
            )
            data["project_experiences"][0]["description"] = (
                "#### 通用子标题\n\n"
                "[链接文字](https://example.com) 和 `代码内容`"
            )
            json_path.write_text(
                json.dumps(data, ensure_ascii=False),
                encoding="utf-8",
            )

            with redirect_stdout(io.StringIO()) as output:
                result = ResumeGenerator(str(json_path), str(output_dir))

            document = Document(result)
            visible_text = _all_document_text(document)
            self.assertNotIn("####", visible_text)
            self.assertNotIn("**", visible_text)
            self.assertNotIn("- 列表", visible_text)
            self.assertIn("C++、C#、char* 和 A*算法", visible_text)
            self.assertNotIn("Markdown残留", output.getvalue())

            list_paragraphs = [
                paragraph
                for paragraph in document.paragraphs
                if paragraph.text.startswith("列表")
            ]
            self.assertEqual(len(list_paragraphs), 2)
            for paragraph in list_paragraphs:
                number_properties = paragraph._p.pPr.find(qn("w:numPr"))
                self.assertIsNotNone(number_properties)

            work_description = document.tables[1].cell(1, 0).paragraphs[0]
            bold_runs = [run.text for run in work_description.runs if run.bold]
            self.assertEqual(bold_runs, ["通用标签："])

            heading = next(
                paragraph
                for paragraph in document.paragraphs
                if paragraph.text == "通用子标题"
            )
            self.assertTrue(all(run.bold for run in heading.runs))
            relationships = document.part.rels.values()
            self.assertTrue(
                any(
                    relationship.reltype.endswith("/hyperlink")
                    and relationship.target_ref == "https://example.com"
                    for relationship in relationships
                )
            )


def _all_document_text(document: Document) -> str:
    """读取正文、表格和文本框文字。"""

    return "\n".join(
        _xml_text(element)
        for element in document._element.body
    )


def _assert_all_declared_fonts_are_yahei(
    test_case: unittest.TestCase,
    docx_path: Path,
) -> None:
    """检查所有显式Word字体声明均为微软雅黑。"""

    namespace = {
        "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
        "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    }

    with ZipFile(docx_path) as archive:
        for name in archive.namelist():
            if not (name.startswith("word/") and name.endswith(".xml")):
                continue

            root = etree.fromstring(archive.read(name))

            for fonts in root.xpath("//w:rFonts", namespaces=namespace):
                for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
                    value = fonts.get(qn(f"w:{attribute}"))

                    if value:
                        test_case.assertEqual(value, FONT_NAME, name)

            for face in root.xpath(
                "//a:rPr/a:latin | //a:rPr/a:ea | //a:rPr/a:cs | "
                "//a:defRPr/a:latin | //a:defRPr/a:ea | //a:defRPr/a:cs",
                namespaces=namespace,
            ):
                test_case.assertEqual(
                    face.get("typeface"),
                    FONT_NAME,
                    name,
                )
