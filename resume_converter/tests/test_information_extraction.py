"""
简历信息提取、证据校验和结构化存储的离线测试。
"""

from __future__ import annotations

import io
import json
import tempfile
import unittest
from contextlib import redirect_stdout
from datetime import date
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

from information_extraction import CheryInformationExtractor, InformationExtractor
from information_extraction import extractor
from information_extraction.normalizer import (
    calculate_work_years,
    normalize_date_range,
    normalize_extraction,
)
from information_extraction.storage import (
    rebuild_summary_workbook,
    save_extracted_json,
)
from utils.processing_records import append_document_record


SAMPLE_MARKDOWN = """
# 张三

性别：男
出生年份：1999
籍贯：湖北武汉

## 专业技能
1. 熟悉 Python、FastAPI
· 掌握 MySQL、Redis

## 工作经历
2024/01-至今 A公司 后端工程师
- 负责使用 Python 开发内部接口

2022/07-2023/12 B公司 开发工程师
- 负责业务系统维护

## 项目经历
项目A｜项目负责人｜2025/01-2025/06
- 负责项目A接口设计和交付

## 教育经历
2022-2026 某大学 本科 计算机科学与技术
""".strip()

CHERY_MARKDOWN = """
# 李四
性别：女
出生年月：1995年6月
联系电话：13800138000
邮箱：lisi@example.com
籍贯：江苏南京

## 工作经历
2020/07-至今 甲公司 测试工程师
负责车载系统测试、缺陷跟踪和版本验收。

## 项目经历
座舱测试项目
项目描述：负责测试方案设计和功能验证。
工作业绩：完成三轮版本验收并推动关键问题闭环。

## 教育经历
2016/09-2020/06 某大学 本科 电子信息工程
""".strip()


def sample_raw_data() -> dict:
    """返回证据可在样例Markdown中逐字核验的模型结果。"""

    return {
        "basic_information": {
            "name": {"value": "张三", "evidence": "# 张三"},
            "gender": {"value": "男", "evidence": "性别：男"},
            "birth_year": {
                "value": "1999",
                "evidence": "出生年份：1999",
            },
            "native_place": {
                "value": "湖北武汉",
                "evidence": "籍贯：湖北武汉",
            },
        },
        "professional_skills": {
            "text": (
                "1. 熟悉 Python、FastAPI\n"
                "· 掌握 MySQL、Redis"
            ),
            "source_type": "explicit",
            "evidence": [
                "1. 熟悉 Python、FastAPI",
                "· 掌握 MySQL、Redis",
            ],
            "note": "来自专业技能栏目",
        },
        "work_experiences": [
            {
                "original_time": "2022/07-2023/12",
                "company_name": "B公司",
                "position_name": "开发工程师",
                "description": "负责业务系统维护",
                "evidence": {
                    "original_time": "2022/07-2023/12",
                    "company_name": "B公司",
                    "position_name": "开发工程师",
                    "description": "负责业务系统维护",
                },
            },
            {
                "original_time": "2024/01-至今",
                "company_name": "A公司",
                "position_name": "后端工程师",
                "description": "负责使用 Python 开发内部接口",
                "evidence": {
                    "original_time": "2024/01-至今",
                    "company_name": "A公司",
                    "position_name": "后端工程师",
                    "description": "负责使用 Python 开发内部接口",
                },
            },
        ],
        "project_experiences": [
            {
                "project_name": "项目A",
                "position_name": "项目负责人",
                "original_time": "2025/01-2025/06",
                "description": "负责项目A接口设计和交付",
                "evidence": {
                    "project_name": "项目A",
                    "position_name": "项目负责人",
                    "original_time": "2025/01-2025/06",
                    "description": "负责项目A接口设计和交付",
                },
            }
        ],
        "education_experiences": [
            {
                "original_time": "2022-2026",
                "school_name": "某大学",
                "degree": "本科",
                "major": "计算机科学与技术",
                "evidence": {
                    "original_time": "2022-2026",
                    "school_name": "某大学",
                    "degree": "本科",
                    "major": "计算机科学与技术",
                },
            }
        ],
        "issues": [],
    }


def sample_chery_raw_data() -> dict:
    """返回含奇瑞扩展字段和受控AI内容的模型结果。"""

    return {
        "basic_information": {
            "name": {"value": "李四", "evidence": "# 李四"},
            "gender": {"value": "女", "evidence": "性别：女"},
            "birth_year": {
                "value": "1995",
                "evidence": "出生年月：1995年6月",
            },
            "birth_date": {
                "value": "1995年6月",
                "evidence": "出生年月：1995年6月",
            },
            "phone": {
                "value": "13800138000",
                "evidence": "联系电话：13800138000",
            },
            "email": {
                "value": "lisi@example.com",
                "evidence": "邮箱：lisi@example.com",
            },
            "native_place": {
                "value": "江苏南京",
                "evidence": "籍贯：江苏南京",
            },
        },
        "self_evaluation": {
            "text": "具备车载系统测试、缺陷跟踪和版本验收经验。",
            "source_type": "generated",
            "evidence": ["负责车载系统测试、缺陷跟踪和版本验收。"],
            "note": "基于工作经历生成",
        },
        "professional_skills": {
            "text": "熟悉车载系统测试、缺陷跟踪和版本验收。",
            "source_type": "generated",
            "evidence": ["负责车载系统测试、缺陷跟踪和版本验收。"],
            "note": "基于工作经历生成",
        },
        "work_experiences": [
            {
                "original_time": "2020/07-至今",
                "company_name": "甲公司",
                "position_name": "测试工程师",
                "description": "负责车载系统测试、缺陷跟踪和版本验收。",
                "evidence": {
                    "original_time": "2020/07-至今",
                    "company_name": "甲公司",
                    "position_name": "测试工程师",
                    "description": "负责车载系统测试、缺陷跟踪和版本验收。",
                },
            }
        ],
        "project_experiences": [
            {
                "project_name": "座舱测试项目",
                "position_name": "",
                "original_time": "",
                "description": "负责测试方案设计和功能验证。",
                "achievement": "完成三轮版本验收并推动关键问题闭环。",
                "evidence": {
                    "project_name": "座舱测试项目",
                    "position_name": "",
                    "original_time": "",
                    "description": "负责测试方案设计和功能验证。",
                    "achievement": "完成三轮版本验收并推动关键问题闭环。",
                },
            }
        ],
        "education_experiences": [
            {
                "original_time": "2016/09-2020/06",
                "school_name": "某大学",
                "degree": "本科",
                "major": "电子信息工程",
                "evidence": {
                    "original_time": "2016/09-2020/06",
                    "school_name": "某大学",
                    "degree": "本科",
                    "major": "电子信息工程",
                },
            }
        ],
        "issues": [],
    }


class NormalizerTests(unittest.TestCase):
    """验证业务规则和防编造校验。"""

    def test_ambiguous_email_marker_forces_user_visible_missing_value(self) -> None:
        """邮箱冲突标记强制清空字段并进入奇瑞待补充流程。"""

        result = normalize_extraction(
            sample_chery_raw_data(),
            CHERY_MARKDOWN + "\n<!-- EMAIL_REQUIRES_MANUAL_REVIEW -->\n",
            calculation_date=date(2026, 8, 1),
            template_tag="奇瑞",
        )

        self.assertEqual(result["basic_information"]["email"], "")
        email_issues = [
            issue
            for issue in result["extraction_issues"]
            if issue["field"] == "basic_information.email"
        ]
        self.assertEqual(len(email_issues), 1)
        self.assertIn("需要人工确认", email_issues[0]["note"])

    def test_prompt_recognizes_products_without_project_section(self) -> None:
        """提示词要求识别工作经历内的明确产品型项目。"""

        self.assertIn("项目不要求出现在独立的", extractor.SYSTEM_PROMPT)
        self.assertIn("主要产品有", extractor.SYSTEM_PROMPT)
        self.assertIn("按原文产品组整体保留", extractor.SYSTEM_PROMPT)
        self.assertIn("完整内容容器", extractor.SYSTEM_PROMPT)
        self.assertIn("不是内容筛选条件", extractor.CHERY_SYSTEM_PROMPT)

    def test_explicit_project_blocks_restore_all_markdown_content(self) -> None:
        """模型只返回概述时，校验层补回项目下的全部原文。"""

        markdown = """
# 王五

## 项目经历
### Gate
**项目描述：** 全球数字资产交易平台。

**应用技术：**
- 使用 Flutter 和 Riverpod 搭建模块
- 通过 Platform Channels 接入原生能力
- 使用 WebSocket 处理实时行情

### 实现细节
- 优化列表帧率和安装包体积

### APDU_SDK
**项目描述：** 安全通信组件。

**应用技术：**
- 支持 SM2、SM3、SM4
- 支持串口、NFC 和 U 口通信

## 教育经历
2016-2020 某大学 本科 软件工程
""".strip()
        raw_data = {
            "basic_information": {},
            "professional_skills": {},
            "work_experiences": [],
            "project_experiences": [
                {
                    "project_name": "Gate",
                    "position_name": "",
                    "original_time": "",
                    "description": "全球数字资产交易平台。",
                    "evidence": {
                        "project_name": "Gate",
                        "position_name": "",
                        "original_time": "",
                        "description": "全球数字资产交易平台。",
                    },
                }
            ],
            "education_experiences": [],
            "issues": [],
        }

        result = normalize_extraction(
            raw_data,
            markdown,
            calculation_date=date(2026, 8, 1),
        )

        self.assertEqual(len(result["project_experiences"]), 2)
        gate = result["project_experiences"][0]
        self.assertFalse(gate["description"].startswith("**项目描述：**"))
        self.assertTrue(gate["description"].startswith("全球数字资产交易平台"))
        self.assertIn("**应用技术：**", gate["description"])
        self.assertIn("Platform Channels", gate["description"])
        self.assertIn("### 实现细节", gate["description"])
        self.assertIn("优化列表帧率和安装包体积", gate["description"])

        recovered = result["project_experiences"][1]
        self.assertEqual(recovered["project_name"], "APDU_SDK")
        self.assertIn("SM2、SM3、SM4", recovered["description"])
        self.assertIn("串口、NFC 和 U 口通信", recovered["description"])
        self.assertTrue(
            any(
                issue["status"] == "recovered"
                for issue in result["extraction_issues"]
            )
        )

    def test_chery_project_recovery_preserves_achievement_and_extra_content(
        self,
    ) -> None:
        """奇瑞工作业绩单列时，其他项目原文仍完整进入项目描述。"""

        markdown = """
## 项目经验
### 座舱测试项目
**项目描述：** 负责测试方案设计和功能验证。

**应用技术：**
- 使用自动化平台执行回归测试
- 跟踪缺陷并完成版本验收

**工作业绩：** 完成三轮版本验收并推动关键问题闭环。
""".strip()
        raw_data = sample_chery_raw_data()
        raw_data["project_experiences"][0]["description"] = (
            "负责测试方案设计和功能验证。"
        )

        result = normalize_extraction(
            raw_data,
            CHERY_MARKDOWN + "\n\n" + markdown,
            calculation_date=date(2026, 8, 1),
            template_tag="奇瑞",
        )

        project = result["project_experiences"][0]
        self.assertIn("使用自动化平台执行回归测试", project["description"])
        self.assertIn("跟踪缺陷并完成版本验收", project["description"])
        self.assertNotIn("**工作业绩：**", project["description"])
        self.assertEqual(
            project["achievement"],
            "完成三轮版本验收并推动关键问题闭环。",
        )

    def test_chery_generated_sections_require_grounded_evidence(self) -> None:
        """奇瑞AI兜底内容必须带原文依据并明确记录人工确认。"""

        self.assertIn("没有时，必须仅依据简历", extractor.CHERY_SYSTEM_PROMPT)
        self.assertIn(
            "只有整份简历不存在任何可作为依据的有效事实时",
            extractor.CHERY_NARRATIVE_FALLBACK_PROMPT,
        )

        result = normalize_extraction(
            sample_chery_raw_data(),
            CHERY_MARKDOWN,
            calculation_date=date(2026, 8, 1),
            template_tag="奇瑞",
        )

        self.assertEqual(
            result["basic_information"]["birth_date"],
            "1995年6月",
        )
        self.assertEqual(result["self_evaluation_source"], "generated")
        self.assertEqual(result["professional_skills_source"], "generated")
        self.assertEqual(
            result["project_experiences"][0]["achievement"],
            "完成三轮版本验收并推动关键问题闭环。",
        )
        generated_fields = {
            issue["field"]
            for issue in result["extraction_issues"]
            if issue["status"] == "generated"
        }
        self.assertEqual(
            generated_fields,
            {"self_evaluation", "professional_skills"},
        )

        ungrounded = sample_chery_raw_data()
        ungrounded["self_evaluation"]["evidence"] = ["原文不存在"]
        rejected = normalize_extraction(
            ungrounded,
            CHERY_MARKDOWN,
            calculation_date=date(2026, 8, 1),
            template_tag="奇瑞",
        )
        self.assertEqual(rejected["self_evaluation"], "")
        self.assertEqual(rejected["self_evaluation_source"], "missing")

        inferred = sample_chery_raw_data()
        inferred["professional_skills"] = {
            "text": "负责车载系统测试、缺陷跟踪和版本验收。",
            "source_type": "inferred",
            "evidence": ["负责车载系统测试、缺陷跟踪和版本验收。"],
            "note": "",
        }
        rejected_inferred = normalize_extraction(
            inferred,
            CHERY_MARKDOWN,
            calculation_date=date(2026, 8, 1),
            template_tag="奇瑞",
        )
        self.assertEqual(rejected_inferred["professional_skills"], "")
        self.assertEqual(
            rejected_inferred["professional_skills_source"],
            "missing",
        )

    def test_product_projects_pass_evidence_validation(self) -> None:
        """产品型项目只要逐字可核验即可进入结构化结果。"""

        markdown = """
2015-2018：深圳航盛扬州研发部
主要产品有：
上汽名爵印度SUV车型车载主机及TBOX产品的开发
青岛四方标准动车组PIS项目开发
""".strip()
        raw_data = {
            "basic_information": {},
            "professional_skills": {},
            "work_experiences": [],
            "project_experiences": [
                {
                    "project_name": (
                        "上汽名爵印度SUV车型车载主机及TBOX产品"
                    ),
                    "position_name": "",
                    "original_time": "2015-2018",
                    "description": (
                        "上汽名爵印度SUV车型车载主机及TBOX产品的开发"
                    ),
                    "evidence": {
                        "project_name": (
                            "上汽名爵印度SUV车型车载主机及TBOX产品"
                        ),
                        "position_name": "",
                        "original_time": "2015-2018",
                        "description": (
                            "上汽名爵印度SUV车型车载主机及TBOX产品的开发"
                        ),
                    },
                },
                {
                    "project_name": "青岛四方标准动车组PIS项目",
                    "position_name": "",
                    "original_time": "2015-2018",
                    "description": "青岛四方标准动车组PIS项目开发",
                    "evidence": {
                        "project_name": "青岛四方标准动车组PIS项目",
                        "position_name": "",
                        "original_time": "2015-2018",
                        "description": "青岛四方标准动车组PIS项目开发",
                    },
                },
            ],
            "education_experiences": [],
            "issues": [],
        }

        result = normalize_extraction(
            raw_data,
            markdown,
            calculation_date=date(2026, 8, 1),
        )

        self.assertEqual(len(result["project_experiences"]), 2)
        self.assertEqual(
            result["project_experiences"][0]["project_name"],
            "上汽名爵印度SUV车型车载主机及TBOX产品",
        )
        self.assertEqual(
            result["project_experiences"][1]["description"],
            "青岛四方标准动车组PIS项目开发",
        )

    def test_education_years_get_default_months_only_for_education(self) -> None:
        """教育年份补9月和6月，工作年份不补月份。"""

        education = normalize_date_range(
            "2022-2026",
            education=True,
        )
        work = normalize_date_range(
            "2022-2026",
            education=False,
        )

        self.assertEqual(education["time"], "2022/09-2026/06")
        self.assertEqual(education["date_status"], "normalized")
        self.assertIn("9月入学", education["normalization_note"])
        self.assertIn("6月毕业", education["normalization_note"])
        self.assertEqual(work["time"], "2022-2026")

    def test_normalization_preserves_skills_and_sorts_work(self) -> None:
        """技能保留原文格式，工作经历按开始时间由近到远。"""

        result = normalize_extraction(
            sample_raw_data(),
            SAMPLE_MARKDOWN,
            calculation_date=date(2026, 8, 1),
        )

        self.assertEqual(
            result["professional_skills"],
            "1. 熟悉 Python、FastAPI\n· 掌握 MySQL、Redis",
        )
        self.assertEqual(
            result["work_experiences"][0]["company_name"],
            "A公司",
        )
        self.assertEqual(
            result["education_experiences"][0]["time"],
            "2022/09-2026/06",
        )
        self.assertEqual(result["work_years"]["value"], 0.0)

    def test_unverifiable_value_is_cleared_and_recorded(self) -> None:
        """模型编造或证据不匹配时清空值并记录问题。"""

        raw_data = sample_raw_data()
        raw_data["basic_information"]["native_place"] = {
            "value": "北京",
            "evidence": "籍贯：北京",
        }
        result = normalize_extraction(
            raw_data,
            SAMPLE_MARKDOWN,
            calculation_date=date(2026, 8, 1),
        )

        self.assertEqual(
            result["basic_information"]["native_place"],
            "",
        )
        self.assertTrue(
            any(
                issue["field"] == "basic_information.native_place"
                for issue in result["extraction_issues"]
            )
        )

    def test_current_residence_is_not_treated_as_native_place(self) -> None:
        """现居地即使有明确地名也不能作为籍贯。"""

        markdown = SAMPLE_MARKDOWN.replace(
            "籍贯：湖北武汉",
            "现居地：湖北武汉",
        )
        raw_data = sample_raw_data()
        raw_data["basic_information"]["native_place"] = {
            "value": "湖北武汉",
            "evidence": "现居地：湖北武汉",
        }
        result = normalize_extraction(
            raw_data,
            markdown,
            calculation_date=date(2026, 8, 1),
        )

        self.assertEqual(
            result["basic_information"]["native_place"],
            "",
        )
        issue = next(
            issue
            for issue in result["extraction_issues"]
            if issue["field"] == "basic_information.native_place"
        )
        self.assertIn("不能用现居地", issue["note"])

    def test_work_years_round_down_to_half_year(self) -> None:
        """毕业后月份按0.5年向下取整。"""

        value = calculate_work_years(
            [
                {
                    "end_date": "2025/01",
                    "original_time": "2021/09-2025/01",
                }
            ],
            calculation_date=date(2026, 8, 1),
        )
        self.assertEqual(value["value"], 1.5)

    def test_ambiguous_degree_is_empty_and_recorded(self) -> None:
        """只写研究生时不擅自判断硕士或博士。"""

        markdown = SAMPLE_MARKDOWN.replace("本科", "研究生")
        raw_data = sample_raw_data()
        raw_data["education_experiences"][0]["degree"] = "研究生"
        raw_data["education_experiences"][0]["evidence"]["degree"] = "研究生"
        result = normalize_extraction(
            raw_data,
            markdown,
            calculation_date=date(2026, 8, 1),
        )

        self.assertEqual(
            result["education_experiences"][0]["degree"],
            "",
        )
        self.assertTrue(
            any(
                issue["field"] == "education_experiences[0].degree"
                for issue in result["extraction_issues"]
            )
        )


class StorageTests(unittest.TestCase):
    """验证JSON和多工作表Excel适合后续Word生成。"""

    def test_json_and_summary_workbook_are_rebuilt(self) -> None:
        """单份JSON可汇总到五张工作表并保留长文本。"""

        with tempfile.TemporaryDirectory(
            dir=Path.cwd()
        ) as temp:
            output_dir = Path(temp)
            normalized = normalize_extraction(
                sample_raw_data(),
                SAMPLE_MARKDOWN,
                calculation_date=date(2026, 8, 1),
            )
            normalized.update(
                {
                    "candidate_id": "candidate-001",
                    "source_file": "D:/input/resume.pdf",
                    "restored_markdown_file": "D:/output/resume_restored.md",
                }
            )
            json_path = output_dir / "resume_extracted.json"
            save_extracted_json(normalized, json_path)
            workbook_path = rebuild_summary_workbook(output_dir)

            workbook = load_workbook(workbook_path, data_only=True)
            self.assertEqual(
                workbook.sheetnames,
                ["候选人汇总", "工作经历", "项目经历", "教育经历", "提取问题"],
            )
            self.assertEqual(
                workbook["候选人汇总"]["F2"].value,
                normalized["professional_skills"],
            )
            self.assertEqual(
                workbook["教育经历"]["C2"].value,
                "2022/09-2026/06",
            )
            self.assertEqual(
                workbook["工作经历"]["F2"].value,
                "A公司",
            )
            workbook.close()


class InformationExtractorTests(unittest.TestCase):
    """验证公开提取接口及处理记录更新，不访问真实LLM。"""

    def test_chery_extractor_uses_template_profile(self) -> None:
        """奇瑞入口保存扩展字段、模板标签和AI来源。"""

        with tempfile.TemporaryDirectory(dir=Path.cwd()) as temp:
            temp_path = Path(temp)
            output_dir = temp_path / "outputs"
            source_path = temp_path / "resume.pdf"
            restored_path = temp_path / "resume_restored.md"
            source_path.write_bytes(b"pdf")
            restored_path.write_text(CHERY_MARKDOWN, encoding="utf-8")

            with (
                patch.dict(
                    extractor.os.environ,
                    {
                        "LLM_API_KEY": "offline-key",
                        "LLM_BASE_URL": "http://127.0.0.1:1/v1",
                        "LLM_MODEL": "offline-model",
                    },
                    clear=False,
                ),
                patch.object(extractor, "load_dotenv"),
                patch.object(
                    extractor,
                    "call_llm_extract_chery",
                    return_value=sample_chery_raw_data(),
                ) as chery_call,
                patch.object(
                    extractor,
                    "_call_llm_complete_chery_narratives",
                ) as fallback_call,
                redirect_stdout(io.StringIO()),
            ):
                result = CheryInformationExtractor(
                    str(source_path),
                    str(restored_path),
                    str(output_dir),
                )

            saved = json.loads(Path(result).read_text(encoding="utf-8"))
            self.assertEqual(saved["template_tag"], "奇瑞")
            self.assertEqual(saved["schema_version"], "1.1")
            self.assertEqual(saved["basic_information"]["phone"], "13800138000")
            self.assertEqual(saved["self_evaluation_source"], "generated")
            chery_call.assert_called_once()
            fallback_call.assert_not_called()

    def test_chery_extractor_retries_missing_generated_sections(self) -> None:
        """奇瑞首轮漏生成长文本时，执行一次受控补全而非直接待补充。"""

        with tempfile.TemporaryDirectory(dir=Path.cwd()) as temp:
            temp_path = Path(temp)
            output_dir = temp_path / "outputs"
            source_path = temp_path / "resume.pdf"
            restored_path = temp_path / "resume_restored.md"
            source_path.write_bytes(b"pdf")
            restored_path.write_text(CHERY_MARKDOWN, encoding="utf-8")

            first_pass = sample_chery_raw_data()
            first_pass["self_evaluation"] = {
                "text": "",
                "source_type": "missing",
                "evidence": [],
                "note": "首轮遗漏",
            }
            first_pass["professional_skills"] = {
                "text": "",
                "source_type": "missing",
                "evidence": [],
                "note": "首轮遗漏",
            }
            first_pass["issues"] = [
                {
                    "field": "self_evaluation",
                    "status": "missing",
                    "evidence": "",
                    "note": "首轮遗漏",
                },
                {
                    "field": "professional_skills",
                    "status": "missing",
                    "evidence": "",
                    "note": "首轮遗漏",
                },
            ]
            generated = sample_chery_raw_data()
            fallback_result = {
                "self_evaluation": generated["self_evaluation"],
                "professional_skills": generated["professional_skills"],
            }

            with (
                patch.dict(
                    extractor.os.environ,
                    {
                        "LLM_API_KEY": "offline-key",
                        "LLM_BASE_URL": "http://127.0.0.1:1/v1",
                        "LLM_MODEL": "offline-model",
                    },
                    clear=False,
                ),
                patch.object(extractor, "load_dotenv"),
                patch.object(
                    extractor,
                    "call_llm_extract_chery",
                    return_value=first_pass,
                ),
                patch.object(
                    extractor,
                    "_call_llm_complete_chery_narratives",
                    return_value=fallback_result,
                ) as fallback_call,
                redirect_stdout(io.StringIO()),
            ):
                result = CheryInformationExtractor(
                    str(source_path),
                    str(restored_path),
                    str(output_dir),
                )

            saved = json.loads(Path(result).read_text(encoding="utf-8"))
            self.assertEqual(saved["self_evaluation_source"], "generated")
            self.assertEqual(saved["professional_skills_source"], "generated")
            fallback_call.assert_called_once()
            stale_missing = {
                issue["field"]
                for issue in saved["extraction_issues"]
                if issue["status"] == "missing"
            }
            self.assertNotIn("self_evaluation", stale_missing)
            self.assertNotIn("professional_skills", stale_missing)

    def test_information_extractor_writes_json_and_updates_tracking(self) -> None:
        """提取成功生成JSON、汇总Excel并更新阶段状态。"""

        with tempfile.TemporaryDirectory(
            dir=Path.cwd()
        ) as temp:
            temp_path = Path(temp)
            output_dir = temp_path / "outputs"
            process_dir = output_dir / "process_data"
            parsing_dir = process_dir / "document_parsing"
            parsing_dir.mkdir(parents=True)
            source_path = temp_path / "resume.pdf"
            restored_path = parsing_dir / "resume_restored.md"
            source_path.write_bytes(b"pdf")
            restored_path.write_text(
                SAMPLE_MARKDOWN,
                encoding="utf-8",
            )
            append_document_record(
                process_dir / "processing_records.xlsx",
                source_path,
                parsing_dir / "resume.md",
                restored_path,
                "成功",
            )

            with (
                patch.dict(
                    extractor.os.environ,
                    {
                        "LLM_API_KEY": "offline-key",
                        "LLM_BASE_URL": "http://127.0.0.1:1/v1",
                        "LLM_MODEL": "offline-model",
                    },
                    clear=False,
                ),
                patch.object(extractor, "load_dotenv"),
                patch.object(
                    extractor,
                    "call_llm_extract",
                    return_value=sample_raw_data(),
                ),
                redirect_stdout(io.StringIO()),
            ):
                result = InformationExtractor(
                    str(source_path),
                    str(restored_path),
                    str(output_dir),
                )

            json_path = Path(result)
            saved = json.loads(
                json_path.read_text(encoding="utf-8")
            )
            self.assertEqual(
                saved["basic_information"]["name"],
                "张三",
            )
            self.assertTrue(
                (
                    process_dir
                    / "information_extraction"
                    / "resume_information.xlsx"
                ).exists()
            )

            tracking = load_workbook(
                process_dir / "processing_records.xlsx",
                data_only=True,
            )
            worksheet = tracking["处理记录"]
            self.assertEqual(worksheet["H2"].value, "成功")
            self.assertEqual(
                worksheet["G2"].value,
                str(json_path.resolve()),
            )
            tracking.close()

    def test_information_extraction_failure_is_logged_and_returns_none(self) -> None:
        """LLM失败不会抛完整堆栈，并更新记录后返回None。"""

        with tempfile.TemporaryDirectory(
            dir=Path.cwd()
        ) as temp:
            temp_path = Path(temp)
            output_dir = temp_path / "outputs"
            process_dir = output_dir / "process_data"
            parsing_dir = process_dir / "document_parsing"
            parsing_dir.mkdir(parents=True)
            source_path = temp_path / "resume.pdf"
            restored_path = parsing_dir / "resume_restored.md"
            source_path.write_bytes(b"pdf")
            restored_path.write_text(SAMPLE_MARKDOWN, encoding="utf-8")
            append_document_record(
                process_dir / "processing_records.xlsx",
                source_path,
                None,
                restored_path,
                "成功",
            )

            with (
                patch.dict(
                    extractor.os.environ,
                    {
                        "LLM_API_KEY": "offline-key",
                        "LLM_BASE_URL": "http://127.0.0.1:1/v1",
                        "LLM_MODEL": "offline-model",
                    },
                    clear=False,
                ),
                patch.object(extractor, "load_dotenv"),
                patch.object(
                    extractor,
                    "call_llm_extract",
                    side_effect=RuntimeError("模拟LLM失败"),
                ),
                redirect_stdout(io.StringIO()),
            ):
                result = InformationExtractor(
                    str(source_path),
                    str(restored_path),
                    str(output_dir),
                )

            self.assertIsNone(result)
            error_file = next(process_dir.glob("error_*.txt"))
            log_content = error_file.read_text(encoding="utf-8")
            self.assertIn("处理阶段：信息提取", log_content)
            self.assertIn("模拟LLM失败", log_content)

            tracking = load_workbook(
                process_dir / "processing_records.xlsx",
                data_only=True,
            )
            worksheet = tracking["处理记录"]
            self.assertEqual(worksheet["H2"].value, "失败")
            self.assertEqual(worksheet["I2"].value, "模拟LLM失败")
            tracking.close()
